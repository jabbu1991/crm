"""Menu in bulk: Excel template / current menu download, Excel (or CSV) upload with a preview, and many photos at once
(the photo's file name is the item name or its S.No)."""
import csv
import io
import json
import os
import re
import secrets

from flask import Response, abort, flash, g, redirect, render_template, request, url_for

from cafe import (FOOD_TYPES, PHOTO_MAX_KB, bp, cafe_name, cafe_view, get_settings, host as H, next_serial, paise_str,
                  to_float, to_paise)

COLUMNS = ["S.No", "Category", "Item name", "Price", "Offer price", "Unit", "Veg / Non-veg", "GST %", "Description",
           "Available", "Barcode"]
UNIT_IN = {"piece": "pc", "pc": "pc", "pcs": "pc", "nos": "pc", "no": "pc", "kg": "kg", "kgs": "kg", "kilo": "kg",
           "litre": "l", "liter": "l", "ltr": "l", "l": "l", "lit": "l"}
UNIT_OUT = {"pc": "Piece", "kg": "Kg", "l": "Litre"}
FOOD_IN = {"veg": "veg", "v": "veg", "non-veg": "nonveg", "nonveg": "nonveg", "non veg": "nonveg", "nv": "nonveg",
           "egg": "egg"}
FOOD_OUT = {"veg": "Veg", "nonveg": "Non-veg", "egg": "Egg"}
MAX_ROWS = 2000


def _org():
    return H.org_id()


def import_dir():
    path = os.path.join(os.path.dirname(H.media_dir()), "imports")
    os.makedirs(path, exist_ok=True)
    return path


def _workbook(rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "Menu"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="C2410C")
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append(r)
    for col, width in zip("ABCDEFGHIJK", (8, 20, 30, 10, 12, 9, 14, 8, 40, 11, 18)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    last = max(MAX_ROWS, len(rows) + 200)
    uv = DataValidation(type="list", formula1='"Piece,Kg,Litre"', allow_blank=True)
    uv.add(f"F2:F{last}")
    dv = DataValidation(type="list", formula1='"Veg,Non-veg,Egg"', allow_blank=True)
    dv.add(f"G2:G{last}")
    av = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    av.add(f"J2:J{last}")
    for v in (uv, dv, av):
        ws.add_data_validation(v)
    for row in ws.iter_rows(min_row=2, min_col=11, max_col=11):
        for c in row:
            c.number_format = "@"  # barcodes stay as text (no 8.9E+12)
    help_ws = wb.create_sheet("How to fill")
    for line in [
        ["How to fill the Menu sheet"],
        [],
        ["Item name and Price are required. Every other column is optional."],
        ["S.No: the item's serial number. Leave empty and it is given automatically."],
        ["Category: e.g. Juices. A new category is created if it does not exist."],
        ["Price / Offer price: numbers only, e.g. 120 or 99.50. Offer price must be lower than Price."],
        ["Unit: Piece, Kg or Litre. Empty = Piece. For Kg / Litre items the price is for 1 kg / 1 litre."],
        ["Barcode: the number printed under the barcode on the packet (optional). Format the column as Text."],
        ["Veg / Non-veg: Veg, Non-veg, Egg or empty."],
        ["GST %: 0, 5, 12, 18 or 28. Empty = the cafe's default GST %."],
        ["Available: Yes or No. Empty = Yes."],
        [],
        ["An item with the same S.No (or the same name) is updated; other rows are added as new items."],
        ["Items that are not in the sheet are not changed or deleted."],
        ["Photos: upload them on the same page. Name each photo file after the item, e.g. 'Chicken Shawarma.jpg',"],
        ["or after its S.No, e.g. '12.jpg'."],
    ]:
        help_ws.append(line)
    help_ws["A1"].font = Font(bold=True, size=14)
    help_ws.column_dimensions["A"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_response(data, name):
    return Response(data, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


def _slug():
    return re.sub(r"[^A-Za-z0-9]+", "-", cafe_name()).strip("-") or "menu"


@bp.route("/cafe/menu/bulk")
@cafe_view("menu_edit")
def menu_bulk():
    count = H.scalar("SELECT COUNT(*) FROM menu_items WHERE org_id = ? AND active = 1", (_org(),))
    no_photo = H.scalar("SELECT COUNT(*) FROM menu_items WHERE org_id = ? AND active = 1 AND photo IS NULL", (_org(),))
    return render_template("cafe/menu_bulk.html", count=count, no_photo=no_photo, columns=COLUMNS)


@bp.route("/cafe/menu/template.xlsx")
@cafe_view("menu_edit")
def menu_template():
    st = get_settings()
    gst = st["default_gst"]
    sample = [[1, "Tea & Coffee", "Kerala Chai", 15, None, "Piece", "Veg", gst, "Fresh milk tea", "Yes", None],
              [2, "Snacks", "Chicken Puffs", 25, None, "Piece", "Non-veg", gst, "", "Yes", None],
              [3, "Bakery", "Mixture", 480, None, "Kg", "Veg", gst, "Price per kg", "Yes", None],
              [4, "Packets", "Choco Brownie", 30, None, "Piece", "Veg", gst, "", "Yes", "8901234567890"]]
    return _xlsx_response(_workbook(sample), "menu-template.xlsx")


@bp.route("/cafe/menu/export.xlsx")
@cafe_view("menu_edit")
def menu_export():
    rows = H.query("""SELECT i.*, c.name cat FROM menu_items i LEFT JOIN menu_categories c ON c.id = i.category_id
                      WHERE i.org_id = ? AND i.active = 1 ORDER BY i.serial, i.name COLLATE NOCASE""", (_org(),))

    def num(p):
        return None if p is None else (p // 100 if p % 100 == 0 else p / 100)
    data = [[r["serial"], r["cat"] or "", r["name"], num(r["price"]), num(r["offer_price"]),
             UNIT_OUT.get(r["unit"] or "pc", "Piece"), FOOD_OUT.get(r["food_type"] or "", ""), r["gst_rate"],
             r["description"] or "", "Yes" if r["available"] else "No", r["barcode"] or None] for r in rows]
    return _xlsx_response(_workbook(data), f"{_slug()}-menu.xlsx")


def _read_rows(upload):
    """Rows (lists of cells) from an .xlsx or .csv upload, header row included."""
    name = (upload.filename or "").lower()
    raw = upload.read()
    if name.endswith(".csv"):
        text = raw.decode("utf-8-sig", errors="replace")
        return [r for r in csv.reader(io.StringIO(text))]
    if not name.endswith((".xlsx", ".xlsm")):
        raise ValueError("Upload the Excel file (.xlsx). Old .xls files: open in Excel and 'Save as' .xlsx first.")
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise ValueError("This file could not be opened as an Excel sheet.")
    ws = wb["Menu"] if "Menu" in wb.sheetnames else wb.worksheets[0]
    out = []
    for row in ws.iter_rows(values_only=True):
        out.append(["" if v is None else v for v in row])
        if len(out) > MAX_ROWS + 1:
            break
    return out


def _cell(row, idx):
    if idx is None or idx >= len(row):
        return ""
    v = row[idx]
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _money(v):
    if v == "":
        return None
    try:
        return int(round(float(str(v).replace(",", "").replace("₹", "").strip()) * 100))
    except ValueError:
        return "bad"


def parse_sheet(rows):
    """Return (lines, errors) where each line is a dict ready to save, with an 'action' (new / update)."""
    if not rows:
        raise ValueError("The sheet is empty.")
    head = [str(h).strip().lower() for h in rows[0]]

    def col(*names):
        for n in names:
            if n in head:
                return head.index(n)
        return None
    idx = {"sno": col("s.no", "sno", "s no", "serial", "serial no"), "cat": col("category"),
           "name": col("item name", "item", "name"), "price": col("price", "rate"),
           "offer": col("offer price", "offer"), "food": col("veg / non-veg", "veg/non-veg", "veg", "type"),
           "gst": col("gst %", "gst", "gst%"), "desc": col("description"), "avail": col("available"),
           "unit": col("unit", "sold by"), "barcode": col("barcode", "bar code", "ean")}
    if idx["name"] is None or idx["price"] is None:
        raise ValueError("The first row must have the column names. 'Item name' and 'Price' are required "
                         "(use the template).")
    oid = _org()
    st = get_settings()
    existing = H.query("SELECT id, name, serial, barcode FROM menu_items WHERE org_id = ? AND active = 1", (oid,))
    by_barcode = {r["barcode"]: r for r in existing if r["barcode"]}
    seen_barcode = set()
    by_serial = {r["serial"]: r for r in existing if r["serial"]}
    by_name = {r["name"].strip().lower(): r for r in existing}
    lines, errors, seen_serial, seen_name = [], [], set(), set()
    for n, row in enumerate(rows[1:], start=2):
        if not any(str(c).strip() for c in row):
            continue
        name = _cell(row, idx["name"])[:80]
        price = _money(_cell(row, idx["price"]))
        offer = _money(_cell(row, idx["offer"]))
        sno_txt = _cell(row, idx["sno"])
        problems = []
        if not name:
            problems.append("item name is empty")
        if price is None or price == "bad" or (isinstance(price, int) and price < 0):
            problems.append("price is missing or not a number")
        if offer == "bad":
            problems.append("offer price is not a number")
        elif isinstance(offer, int) and isinstance(price, int) and offer >= price:
            problems.append("offer price must be lower than the price")
        sno = None
        if sno_txt:
            if re.fullmatch(r"\d{1,5}", sno_txt) and int(sno_txt) > 0:
                sno = int(sno_txt)
            else:
                problems.append("S.No must be a whole number")
        food_txt = _cell(row, idx["food"]).lower()
        food = FOOD_IN.get(food_txt, "") if food_txt else ""
        if food_txt and food_txt not in FOOD_IN:
            problems.append("Veg / Non-veg must be Veg, Non-veg or Egg")
        gst_txt = _cell(row, idx["gst"]).replace("%", "")
        gst = to_float(gst_txt, None, 0, 28) if gst_txt else st["default_gst"]
        if gst is None:
            problems.append("GST % is not a number")
        unit_txt = _cell(row, idx["unit"]).lower().strip(". ")
        unit = UNIT_IN.get(unit_txt, "pc") if unit_txt else "pc"
        if unit_txt and unit_txt not in UNIT_IN:
            problems.append("Unit must be Piece, Kg or Litre")
        barcode = re.sub(r"\s+", "", _cell(row, idx["barcode"]))[:40] or None
        if barcode and barcode in seen_barcode:
            problems.append(f"barcode {barcode} appears twice in the sheet")
        avail_txt = _cell(row, idx["avail"]).lower()
        available = 0 if avail_txt in ("no", "n", "0", "off", "false") else 1
        key = name.lower()
        if sno and sno in seen_serial:
            problems.append(f"S.No {sno} appears twice in the sheet")
        if key and key in seen_name:
            problems.append("the same item name appears twice in the sheet")
        target = by_serial.get(sno) if sno else None
        if target is None:
            target = by_name.get(key)
        if sno and target is None and sno in by_serial:
            target = by_serial[sno]
        if sno and target is not None and sno in by_serial and by_serial[sno]["id"] != target["id"]:
            problems.append(f"S.No {sno} already belongs to '{by_serial[sno]['name']}'")
        if barcode and barcode in by_barcode and (target is None or by_barcode[barcode]["id"] != target["id"]):
            problems.append(f"barcode {barcode} already belongs to '{by_barcode[barcode]['name']}'")
        if problems:
            errors.append({"row": n, "name": name, "problems": problems})
            continue
        if barcode:
            seen_barcode.add(barcode)
        seen_serial.add(sno)
        seen_name.add(key)
        lines.append({"row": n, "action": "update" if target else "new", "id": target["id"] if target else None,
                      "serial": sno, "category": _cell(row, idx["cat"])[:60], "name": name, "price": price,
                      "offer_price": offer, "food_type": food, "gst_rate": gst,
                      "description": _cell(row, idx["desc"])[:300], "available": available, "unit": unit,
                      "barcode": barcode})
    return lines, errors


@bp.route("/cafe/menu/import", methods=["POST"])
@cafe_view("menu_edit")
def menu_import():
    f = request.files.get("sheet")
    if not f or not f.filename:
        flash("Choose the Excel file to upload.", "error")
        return redirect(url_for("cafe.menu_bulk"))
    try:
        lines, errors = parse_sheet(_read_rows(f))
    except ValueError as e:
        flash(str(e), "error")
        return redirect(url_for("cafe.menu_bulk"))
    if not lines and not errors:
        flash("No items found in the sheet.", "error")
        return redirect(url_for("cafe.menu_bulk"))
    tok = secrets.token_hex(12)
    with open(os.path.join(import_dir(), tok + ".json"), "w") as out:
        json.dump({"org": _org(), "user": g.user["id"], "lines": lines}, out)
    new = sum(1 for l in lines if l["action"] == "new")
    return render_template("cafe/menu_import_preview.html", lines=lines, errors=errors, tok=tok, new=new,
                           upd=len(lines) - new, paise_str=paise_str, FOOD_OUT=FOOD_OUT)


@bp.route("/cafe/menu/import/confirm", methods=["POST"])
@cafe_view("menu_edit")
def menu_import_confirm():
    tok = request.form.get("tok") or ""
    if not re.fullmatch(r"[0-9a-f]{24}", tok):
        abort(400)
    path = os.path.join(import_dir(), tok + ".json")
    try:
        with open(path) as fh:
            data = json.load(fh)
    except OSError:
        flash("This preview has expired. Upload the sheet again.", "error")
        return redirect(url_for("cafe.menu_bulk"))
    if data.get("org") != _org():
        abort(404)
    oid, ts = _org(), H.now()
    db = H.get_db()
    cats = {r["name"].lower(): r["id"] for r in
            db.execute("SELECT id, name FROM menu_categories WHERE org_id = ? AND active = 1", (oid,)).fetchall()}
    next_sort = (db.execute("SELECT COALESCE(MAX(sort), 0) FROM menu_categories WHERE org_id = ?",
                            (oid,)).fetchone()[0] or 0) + 1
    serial = next_serial(oid)
    taken = {r[0] for r in db.execute("SELECT serial FROM menu_items WHERE org_id = ? AND active = 1 AND serial IS NOT NULL",
                                      (oid,)).fetchall()}
    added = updated = 0
    try:
        for l in data["lines"]:
            cid = None
            if l["category"]:
                cid = cats.get(l["category"].lower())
                if cid is None:
                    cid = db.execute("INSERT INTO menu_categories (org_id, name, sort, created_at) VALUES (?,?,?,?)",
                                     (oid, l["category"], next_sort, ts)).lastrowid
                    cats[l["category"].lower()] = cid
                    next_sort += 1
            vals = (l["name"], l["description"] or None, cid, l["price"], l["offer_price"], l["food_type"] or None,
                    l["gst_rate"], l["available"], l.get("unit") or "pc", l.get("barcode"))
            if l["action"] == "update":
                row = db.execute("SELECT id, serial FROM menu_items WHERE id = ? AND org_id = ? AND active = 1",
                                 (l["id"], oid)).fetchone()
                if not row:
                    continue
                sno = l["serial"] or row["serial"]
                db.execute("""UPDATE menu_items SET name=?, description=?, category_id=?, price=?, offer_price=?,
                              food_type=?, gst_rate=?, available=?, unit=?, barcode=?, serial=?, updated_at=?
                              WHERE id=?""",
                           vals + (sno, ts, row["id"]))
                updated += 1
            else:
                sno = l["serial"] if l["serial"] and l["serial"] not in taken else None
                if sno is None:
                    while serial in taken:
                        serial += 1
                    sno = serial
                taken.add(sno)
                db.execute("""INSERT INTO menu_items (org_id, name, description, category_id, price, offer_price,
                              food_type, gst_rate, available, unit, barcode, serial, created_by, created_at, updated_at)
                              VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", (oid,) + vals + (sno, g.user["id"], ts, ts))
                added += 1
        db.commit()
    except Exception:
        db.rollback()
        raise
    try:
        os.remove(path)
    except OSError:
        pass
    flash(f"Menu imported: {added} new items added, {updated} items updated.", "success")
    return redirect(url_for("cafe.menu"))


def _norm(s):
    return re.sub(r"[^a-z0-9]+", "", (s or "").lower())


@bp.route("/cafe/menu/photos", methods=["POST"])
@cafe_view("menu_edit")
def menu_photos():
    files = [f for f in request.files.getlist("photos") if f and f.filename]
    if not files:
        flash("Choose the photos to upload.", "error")
        return redirect(url_for("cafe.menu_bulk"))
    items = H.query("SELECT id, name, serial, photo FROM menu_items WHERE org_id = ? AND active = 1", (_org(),))
    by_name = {_norm(i["name"]): i for i in items}
    by_serial = {str(i["serial"]): i for i in items if i["serial"]}
    matched, unmatched, bad = [], [], []
    for f in files[:300]:
        stem = os.path.splitext(os.path.basename(f.filename))[0].strip()
        item = by_serial.get(stem.lstrip("0") or stem) if stem.isdigit() else None
        item = item or by_name.get(_norm(stem))
        if not item:
            unmatched.append(f.filename)
            continue
        name = H.save_image(f, box=(600, 600), max_kb=PHOTO_MAX_KB, fit=True)
        if not name:
            bad.append(f.filename)
            continue
        H.remove_media(item["photo"])
        H.execute("UPDATE menu_items SET photo = ?, updated_at = ? WHERE id = ?", (name, H.now(), item["id"]))
        matched.append(item["name"])
    if matched:
        flash(f"{len(matched)} photos added: {', '.join(matched[:15])}{' …' if len(matched) > 15 else ''}", "success")
    if unmatched:
        flash(f"{len(unmatched)} photos did not match any item name or S.No: {', '.join(unmatched[:15])}"
              f"{' …' if len(unmatched) > 15 else ''}. Rename them to the item name (e.g. 'Kerala Chai.jpg') and "
              "upload again.", "warning")
    if bad:
        flash(f"{len(bad)} files were not photos: {', '.join(bad[:10])}", "warning")
    return redirect(url_for("cafe.menu_bulk"))
