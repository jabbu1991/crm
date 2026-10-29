"""Production: the printable production sheet for the maker, the customer's private tracking / warranty page,
Excel export of all production data, and bulk import of deliveries made before the software."""
import io
import json
import os
import re
import secrets
from datetime import timedelta

from flask import Response, abort, flash, g, redirect, render_template, request, url_for

from production import (MARKS, bp, can, clean, create_order, customer_by_code, details_pending, get_order,
                        get_settings, has_prod, host as H, int_in, marks_of, order_days, order_rows, phase,
                        prod_view, schedule, set_stages, to_date, today, warranty_info)

MAX_ROWS = 3000


# ---------------------------------------------------------------- production sheet (A4, for the maker)

@bp.route("/production/orders/<int:order_id>/sheet")
@prod_view("view")
def order_sheet(order_id):
    o = get_order(None, order_id)
    brand = H.org_brand() or {"name": g.org["name"], "logo": None}
    show_amount = can("maker")
    return render_template("production/sheet.html", o=o, brand=brand, st=get_settings(), show_amount=show_amount,
                           org=g.org, days=order_days(o))


# ---------------------------------------------------------------- customer page (no login)

@bp.route("/t/<code>/<tok>")
def public_track(code, tok):
    org = H.query("SELECT * FROM organizations WHERE LOWER(code) = ? AND status = 'Active'", (code.lower(),), one=True)
    if not org or not has_prod(org) or not re.fullmatch(r"[A-Za-z0-9_-]{8,40}", tok or ""):
        abort(404)
    st = get_settings(org["id"])
    if not st["customer_link"]:
        abort(404)
    o = get_order(org["id"], H.scalar("SELECT id FROM prod_orders WHERE org_id = ? AND token = ?", (org["id"], tok))
                  or abort(404))
    g.org = org  # for the helpers below (read-only page)
    stages = set_stages(o["set_id"])
    log = H.query("SELECT stage_name, created_at FROM prod_stage_log WHERE order_id = ? ORDER BY id", (o["id"],))
    reached = {}
    for r in log:
        reached[r["stage_name"]] = r["created_at"]
    parts = H.query("SELECT * FROM prod_parts WHERE order_id = ? ORDER BY sort, id", (o["id"],)) if st["parts"] else []
    ph = phase(o)
    sched = schedule(org["id"]).get(o["id"]) if ph in ("queue", "working", "ready") else None
    ledger, value = [], 0
    if st["referral"]:
        ledger = H.query("""SELECT p.change, p.kind, p.amount, p.note, p.created_at, p.pay_date FROM prod_points p
                            WHERE p.customer_id = ? ORDER BY p.id DESC LIMIT 50""", (o["customer_id"],))
        value = (o["cust_points"] or 0) * int(st["point_value"] or 0)
    brand = H.org_brand(org) or {"name": org["name"], "logo": None, "palette": None}
    share = ""
    if st["referral"]:
        from urllib.parse import quote
        share = "https://wa.me/?text=" + quote(
            f"I bought from {brand['name']}. Order with my code {o['unit_code']} 🎁 — call {org['phone'] or ''}".strip(" —"))
    return render_template("production/public.html", o=o, org=org, st=st, stages=stages, reached=reached,
                           w=warranty_info(o, parts) if st["warranty"] else None, ph=ph, sched=sched, ledger=ledger,
                           value=value, brand=brand, share=share, cancelled=o["status"] != "active",
                           delivered=ph == "delivered")


# ---------------------------------------------------------------- Excel helpers

def _xlsx_response(wb, name):
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


def _sheet(wb, title, head, rows, widths=None, first=False):
    from openpyxl.styles import Font, PatternFill
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="EA580C")
    for r in rows:
        ws.append([_safe(v) for v in r])
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"
    return ws


def _safe(v):
    """Text that starts like a formula is written as plain text."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") and not re.fullmatch(r"-?\d+(\.\d+)?", v):
        return "'" + v
    return v


def _rs(p):
    return None if p is None else (p // 100 if p % 100 == 0 else p / 100)


def _slug():
    return re.sub(r"[^A-Za-z0-9]+", "-", g.org["name"]).strip("-") or "production"


# ---------------------------------------------------------------- Import / Export page

@bp.route("/production/data")
@prod_view()
def data():
    if not (can("import") or can("export")):
        abort(403)
    oid = H.org_id()
    stages = [r[0] for r in H.query("""SELECT name FROM prod_stages WHERE org_id = ? AND active = 1 GROUP BY name
                                       ORDER BY MIN(sort)""", (oid,))]
    products = [r[0] for r in H.query("SELECT DISTINCT product_name FROM prod_orders WHERE org_id = ? AND"
                                      " product_name != '' ORDER BY product_name", (oid,))]
    return render_template("production/data.html", stages=stages, products=products, columns=IMPORT_COLUMNS,
                           count=H.scalar("SELECT COUNT(*) FROM prod_orders WHERE org_id = ?", (oid,)))


@bp.route("/production/export.xlsx")
@prod_view("export")
def export_all():
    from openpyxl import Workbook
    oid = H.org_id()
    a = request.args
    where, args = [], []
    if H.valid_date(a.get("from") or ""):
        where.append("o.order_date >= ?")
        args.append(a.get("from"))
    if H.valid_date(a.get("to") or ""):
        where.append("o.order_date <= ?")
        args.append(a.get("to"))
    if a.get("stage"):
        where.append("s.name = ?")
        args.append(a.get("stage"))
    if a.get("product"):
        where.append("o.product_name = ?")
        args.append(a.get("product"))
    if not a.get("cancelled"):
        where.append("o.status = 'active'")
    rows = order_rows(oid, ("AND " + " AND ".join(where)) if where else "", args, order_by="o.order_no")
    ids = [o["id"] for o in rows]
    sched = schedule(oid)
    maker_ok = can("maker")
    adv = {}
    for r in H.query("SELECT order_id, SUM(amount) t FROM prod_advances WHERE org_id = ? GROUP BY order_id", (oid,)):
        adv[r["order_id"]] = r["t"]
    parts_by = {}
    for p in H.query("SELECT * FROM prod_parts WHERE org_id = ? AND order_id IS NOT NULL ORDER BY order_id, sort, id",
                     (oid,)):
        parts_by.setdefault(p["order_id"], []).append(p)
    wb = Workbook()
    head = ["Order no", "Order date", "Customer", "Phone", "WhatsApp", "Address", "Unit code", "Product", "Category",
            "Machine type", "Power source", "Weight", "Quantity", "Unit price", "Advance received", "Details",
            "Stage", "Status", "Production days", "MFG start", "Expected finish / dispatch", "Late",
            "Production finished", "First advance", "Moved to top (reason)",
            "Manufacturing in-charge"] + (["Maker amount"] if maker_ok else []) + [
            "GST bill no", "Bill date", "Warranty months", "Warranty ends", "Warranty days left", "Delivered date",
            "Order taken by", "Approved by", "Referral code used", "Imported (old delivery)", "Cancel reason", "Note",
            "Customer link"]
    data_rows = []
    for o in rows:
        w = warranty_info(o)["main"]
        sc = sched.get(o["id"], {})
        link = url_for("prod.public_track", code=g.org["code"].lower(), tok=o["token"], _external=True)
        data_rows.append([o["order_no"], o["order_date"], o["cust_name"], o["cust_phone"], o["cust_wa"],
                          H.scalar("SELECT address FROM prod_customers WHERE id = ?", (o["customer_id"],)),
                          o["unit_code"], o["product_name"], o["category"], o["machine_type"], o["power_source"],
                          o["weight"], o["quantity"], _rs(o["unit_price"]), _rs(adv.get(o["id"])), o["details"],
                          o["stage_name"], o["status"].title(), o["prod_days"], o["mfg_start"],
                          sc["finish"].isoformat() if sc.get("finish") else (o["exp_finish"] or ""),
                          "Yes" if sc.get("late") else "", (o["finished_at"] or "")[:10], o["first_adv"],
                          o["pin_reason"] if o["pinned_at"] else "", o["maker_name"]]
                         + ([_rs(o["maker_amount"])] if maker_ok else [])
                         + [o["bill_no"], o["bill_date"], o["warranty_months"], w["end"].isoformat() if w else "",
                            w["left"] if w else "", o["delivered_date"], o["taken_by_name"], o["approved_by"],
                            o["referral_code"], "Yes" if o["imported"] else "", o["cancel_reason"], o["note"], link])
    _sheet(wb, "Orders", head, data_rows, [9, 11, 26, 14, 14, 24, 10, 26, 14, 12, 12, 10, 8, 11, 11, 40, 18, 10, 10,
                                            11, 14, 6, 18], first=True)
    if ids:
        marks = ",".join("?" * len(ids))
        hist = H.query(f"""SELECT o.order_no, l.created_at, l.from_name, l.stage_name, l.note, u.name FROM prod_stage_log l
                           JOIN prod_orders o ON o.id = l.order_id LEFT JOIN users u ON u.id = l.user_id
                           WHERE l.order_id IN ({marks}) ORDER BY o.order_no, l.id""", ids)
    else:
        hist = []
    _sheet(wb, "Stage history", ["Order no", "When", "From", "To", "Note", "By"],
           [[h[0], h[1], h[2], h[3], h[4], h[5]] for h in hist], [9, 20, 18, 18, 40, 18])
    prow = []
    for o in rows:
        info = warranty_info(o, parts_by.get(o["id"], []))
        for x in info["parts"]:
            p = x["p"]
            prow.append([o["order_no"], o["bill_no"], p["name"], p["company"], p["months"], p["serial"],
                         x["end"].isoformat() if x["end"] else "", x["left"] if x["left"] is not None else ""])
    _sheet(wb, "Parts warranty", ["Order no", "GST bill no", "Part", "Company", "Warranty months", "Serial no",
                                  "Warranty ends", "Days left"], prow, [9, 14, 24, 22, 10, 18, 14, 10])
    st = get_settings()
    cust = H.query("""SELECT c.*, (SELECT COUNT(*) FROM prod_orders o WHERE o.customer_id = c.id) n FROM prod_customers c
                      WHERE c.org_id = ? ORDER BY c.name COLLATE NOCASE""", (oid,))
    _sheet(wb, "Customers", ["Name", "Phone", "WhatsApp", "Address", "Unit code", "Orders", "Points", "Points value"],
           [[c["name"], c["phone"], c["whatsapp"], c["address"], c["unit_code"], c["n"], c["points"],
             _rs(c["points"] * int(st["point_value"] or 0))] for c in cust], [26, 14, 14, 24, 10, 8, 8, 12])
    pts = H.query("""SELECT p.*, c.name, c.unit_code, o.order_no, u.name uname FROM prod_points p
                     JOIN prod_customers c ON c.id = p.customer_id LEFT JOIN prod_orders o ON o.id = p.order_id
                     LEFT JOIN users u ON u.id = p.created_by WHERE p.org_id = ? ORDER BY p.id""", (oid,))
    _sheet(wb, "Points history", ["When", "Customer", "Unit code", "Points", "Type", "Order no", "Amount paid",
                                  "Paid by", "Paid on", "Note", "By"],
           [[p["created_at"], p["name"], p["unit_code"], p["change"], p["kind"].title(), p["order_no"], _rs(p["amount"]),
             p["pay_mode"], p["pay_date"], p["note"], p["uname"]] for p in pts], [20, 24, 10, 8, 10, 9, 12, 12, 12, 30, 16])
    if maker_ok:
        mk = H.query("""SELECT m.*, (SELECT COUNT(*) FROM prod_orders o WHERE o.maker_id = m.id) n,
                          (SELECT COALESCE(SUM(maker_amount), 0) FROM prod_orders o WHERE o.maker_id = m.id
                             AND o.status = 'active') t FROM prod_makers m WHERE m.org_id = ? ORDER BY m.name""", (oid,))
        _sheet(wb, "Makers", ["Name", "Phone", "Work", "Orders", "Total maker amount", "Active", "Note"],
               [[m["name"], m["phone"], m["work"], m["n"], _rs(m["t"]), "Yes" if m["active"] else "No", m["note"]]
                for m in mk], [22, 14, 18, 8, 14, 7, 30])
        mrows = [[o["order_no"], o["maker_name"], o["product_name"], o["cust_name"], o["mfg_start"], o["exp_finish"],
                  _rs(o["maker_amount"]), o["stage_name"]] for o in rows if o["maker_id"]]
        _sheet(wb, "Maker amounts", ["Order no", "Maker", "Product", "Customer", "MFG start", "Expected finish",
                                     "Maker amount", "Stage"], mrows, [9, 20, 26, 24, 11, 14, 12, 18])
    H.audit("Production data exported (Excel)", f"{len(rows)} orders")
    return _xlsx_response(wb, f"{_slug()}-production-{today()}.xlsx")


# ---------------------------------------------------------------- import old deliveries

IMPORT_COLUMNS = ["Customer name", "Phone", "WhatsApp", "Address / Place", "Product", "Category", "Machine type",
                  "Power source", "Weight", "Quantity", "Unit price", "Order date", "Delivered date", "GST bill no",
                  "Bill date", "Warranty months", "Maker", "Note"]
PART_COLUMNS = ["GST bill no", "Part name", "Company", "Warranty months", "Serial no"]


def import_dir():
    path = os.path.join(os.path.dirname(H.media_dir()), "imports")
    os.makedirs(path, exist_ok=True)
    return path


@bp.route("/production/import/template.xlsx")
@prod_view("import")
def import_template():
    from openpyxl import Workbook
    wb = Workbook()
    ws = _sheet(wb, "Deliveries", IMPORT_COLUMNS,
                [["Shreesalasar Agro Pvt Ltd", "9876543210", "9876543210", "West Bengal", "Vertical Electric with Stand",
                  "Crompton 7.5", "New", "Electric", "", 1, 14000, "2025-03-02", "2025-03-20", "GST/2025/101",
                  "2025-03-18", 12, "Babu", "Example row — delete it"]],
                [26, 14, 14, 22, 28, 14, 12, 12, 9, 8, 11, 12, 13, 14, 12, 10, 14, 26], first=True)
    for row in ws.iter_rows(min_row=2, max_row=MAX_ROWS, min_col=2, max_col=3):
        for c in row:
            c.number_format = "@"
    _sheet(wb, "Parts", PART_COLUMNS, [["GST/2025/101", "Motor", "Crompton", 24, "CR123456"]], [14, 22, 20, 10, 18])
    help_ws = wb.create_sheet("How to fill")
    for line in [["How to fill"], [],
                 ["Deliveries sheet: one row = one delivered product (one order)."],
                 ["Required: Customer name, Phone (or WhatsApp), Product, Bill date. Everything else is optional."],
                 ["Dates: YYYY-MM-DD (e.g. 2025-03-18) or DD-MM-YYYY. Unit price: numbers only."],
                 ["Warranty months empty = the product's warranty (Production → Products) or the default."],
                 ["The same phone number = the same customer (one customer, many orders)."],
                 ["A GST bill number that is already in the software is skipped (no duplicates)."],
                 ["Parts sheet (optional): parts made by other companies, matched to the delivery by GST bill no."],
                 ["Imported orders go straight to Delivered. They give no referral points."]]:
        help_ws.append(line)
    help_ws.column_dimensions["A"].width = 100
    return _xlsx_response(wb, "old-deliveries-template.xlsx")


def _rows(ws):
    out = []
    for row in ws.iter_rows(values_only=True):
        out.append(["" if v is None else v for v in row])
        if len(out) > MAX_ROWS + 1:
            break
    return out


def _cell(row, idx):
    if idx is None or idx >= len(row):
        return ""
    v = row[idx]
    if hasattr(v, "date") and callable(getattr(v, "date")):
        return v.date().isoformat()
    if hasattr(v, "isoformat") and not isinstance(v, str):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _date(text):
    text = (text or "").strip()[:10]
    if not text:
        return ""
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        return H.valid_date(text)
    m = re.fullmatch(r"(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})", text)
    if m:
        return H.valid_date(f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}")
    return None


def _col(head, *names):
    for n in names:
        if n in head:
            return head.index(n)
    return None


def parse_import(raw, filename):
    name = (filename or "").lower()
    if not name.endswith((".xlsx", ".xlsm")):
        raise ValueError("Upload the Excel file (.xlsx) — use the template.")
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise ValueError("This file could not be opened as an Excel sheet.")
    ws = wb["Deliveries"] if "Deliveries" in wb.sheetnames else wb.worksheets[0]
    rows = _rows(ws)
    if not rows:
        raise ValueError("The sheet is empty.")
    head = [str(h).strip().lower() for h in rows[0]]
    idx = {k: _col(head, *names) for k, names in {
        "name": ("customer name", "customer", "name"), "phone": ("phone", "mobile"), "wa": ("whatsapp",),
        "address": ("address / place", "address", "place"), "product": ("product", "product name"),
        "category": ("category", "product category"), "machine": ("machine type",), "power": ("power source",),
        "weight": ("weight",), "qty": ("quantity", "qty"), "price": ("unit price", "price"),
        "order_date": ("order date",), "delivered": ("delivered date", "delivery date"),
        "bill_no": ("gst bill no", "bill no", "gst bill number"), "bill_date": ("bill date",),
        "warranty": ("warranty months", "warranty"), "maker": ("maker", "manufacturing in-charge"),
        "note": ("note", "notes")}.items()}
    if idx["name"] is None or idx["bill_date"] is None or idx["product"] is None:
        raise ValueError("The first row must have the column names (Customer name, Product, Bill date…). "
                         "Use the template.")
    parts = {}
    if "Parts" in wb.sheetnames:
        prow = _rows(wb["Parts"])
        if prow:
            ph = [str(h).strip().lower() for h in prow[0]]
            pi = {"bill": _col(ph, "gst bill no", "bill no"), "name": _col(ph, "part name", "part"),
                  "company": _col(ph, "company"), "months": _col(ph, "warranty months", "warranty"),
                  "serial": _col(ph, "serial no", "serial")}
            for r in prow[1:]:
                b, n = _cell(r, pi["bill"]), _cell(r, pi["name"])
                if b and n:
                    parts.setdefault(b, []).append({"name": n[:120], "company": _cell(r, pi["company"])[:120],
                                                    "months": int_in(_cell(r, pi["months"]), None, 0, 600),
                                                    "serial": _cell(r, pi["serial"])[:80]})
    oid = H.org_id()
    existing_bills = {r[0] for r in H.query("SELECT bill_no FROM prod_orders WHERE org_id = ? AND bill_no IS NOT NULL",
                                            (oid,))}
    products = {r["name"].lower(): r for r in H.query("SELECT id, name, category FROM products WHERE org_id = ?", (oid,))}
    lines, errors, skipped, seen = [], [], [], set()
    for n, row in enumerate(rows[1:], start=2):
        if not any(str(c).strip() for c in row):
            continue
        note = _cell(row, idx["note"])
        if "example row" in note.lower():
            continue
        name = _cell(row, idx["name"])[:120]
        phone, wa = _cell(row, idx["phone"])[:40], _cell(row, idx["wa"])[:40]
        product = _cell(row, idx["product"])[:120]
        bill_no = _cell(row, idx["bill_no"])[:60]
        problems = []
        if not name:
            problems.append("customer name is empty")
        if not (phone or wa):
            problems.append("phone / WhatsApp is empty")
        if not product:
            problems.append("product is empty")
        dates = {}
        for k, label in (("order_date", "order date"), ("delivered", "delivered date"), ("bill_date", "bill date")):
            v = _date(_cell(row, idx[k]))
            if v is None:
                problems.append(f"{label} is not a date (use YYYY-MM-DD)")
            dates[k] = v or ""
        if not dates["bill_date"] and "bill date is not a date (use YYYY-MM-DD)" not in problems:
            problems.append("bill date is empty (warranty starts from it)")
        price_txt = _cell(row, idx["price"]).replace(",", "").replace("₹", "")
        price = None
        if price_txt:
            try:
                price = int(round(float(price_txt) * 100))
            except ValueError:
                problems.append("unit price is not a number")
        wtxt = _cell(row, idx["warranty"])
        warranty = int_in(wtxt, None, 0, 600) if wtxt else None
        if wtxt and warranty is None:
            problems.append("warranty months is not a number")
        if bill_no and bill_no in seen:
            problems.append(f"GST bill no {bill_no} appears twice in the sheet")
        if problems:
            errors.append({"row": n, "name": name, "problems": problems})
            continue
        if bill_no and bill_no in existing_bills:
            skipped.append({"row": n, "name": name, "bill_no": bill_no})
            continue
        if bill_no:
            seen.add(bill_no)
        prod = products.get(product.lower())
        lines.append({"row": n, "name": name, "phone": phone, "whatsapp": wa or phone,
                      "address": _cell(row, idx["address"])[:200], "product_id": prod["id"] if prod else None,
                      "product_name": prod["name"] if prod else product,
                      "category": _cell(row, idx["category"])[:80] or (prod["category"] if prod else ""),
                      "machine_type": _cell(row, idx["machine"])[:60], "power_source": _cell(row, idx["power"])[:60],
                      "weight": _cell(row, idx["weight"])[:40], "quantity": int_in(_cell(row, idx["qty"]), 1, 1, 100000),
                      "unit_price": price, "order_date": dates["order_date"] or dates["bill_date"],
                      "delivered_date": dates["delivered"] or dates["bill_date"], "bill_no": bill_no,
                      "bill_date": dates["bill_date"], "warranty_months": warranty, "maker": _cell(row, idx["maker"])[:120],
                      "note": note[:500], "parts": parts.get(bill_no, []) if bill_no else []})
    return lines, errors, skipped


@bp.route("/production/import", methods=["POST"])
@prod_view("import")
def import_upload():
    f = request.files.get("sheet")
    if not f or not f.filename:
        flash("Choose the Excel file to upload.", "error")
        return redirect(url_for("prod.data"))
    try:
        lines, errors, skipped = parse_import(f.read(), f.filename)
    except ValueError as e:
        flash(str(e), "error")
        return redirect(url_for("prod.data"))
    if not lines and not errors and not skipped:
        flash("No rows found in the sheet.", "error")
        return redirect(url_for("prod.data"))
    tok = secrets.token_hex(12)
    with open(os.path.join(import_dir(), "prod-" + tok + ".json"), "w") as out:
        json.dump({"org": H.org_id(), "user": g.user["id"], "lines": lines}, out)
    customers = len({H.phone_key(l["phone"] or l["whatsapp"]) for l in lines})
    return render_template("production/import_preview.html", lines=lines, errors=errors, skipped=skipped, tok=tok,
                           customers=customers)


@bp.route("/production/import/confirm", methods=["POST"])
@prod_view("import")
def import_confirm():
    tok = request.form.get("tok") or ""
    if not re.fullmatch(r"[0-9a-f]{24}", tok):
        abort(400)
    path = os.path.join(import_dir(), "prod-" + tok + ".json")
    try:
        with open(path) as fh:
            saved = json.load(fh)
    except (OSError, ValueError):
        flash("This upload has expired. Please upload the file again.", "error")
        return redirect(url_for("prod.data"))
    if saved.get("org") != H.org_id() or saved.get("user") != g.user["id"]:
        abort(403)
    db = H.get_db()
    ts = H.now()
    oid = H.org_id()
    done = 0
    bills = {r[0] for r in H.query("SELECT bill_no FROM prod_orders WHERE org_id = ? AND bill_no IS NOT NULL", (oid,))}
    makers = {}
    for l in saved["lines"]:
        if l["bill_no"] and l["bill_no"] in bills:
            continue  # added meanwhile
        mid = None
        if l.get("maker"):
            key = l["maker"].lower()
            if key not in makers:
                row = db.execute("SELECT id FROM prod_makers WHERE org_id = ? AND name = ? COLLATE NOCASE",
                                 (oid, l["maker"])).fetchone()
                makers[key] = row[0] if row else db.execute(
                    "INSERT INTO prod_makers (org_id, name, created_at) VALUES (?,?,?)", (oid, l["maker"], ts)).lastrowid
            mid = makers[key]
        d = dict(l, maker_id=mid, referral_code="")
        new_id = create_order(db, oid, d, imported=True, ts=ts, note="Imported (old delivery)")
        if l["parts"]:
            db.execute("DELETE FROM prod_parts WHERE order_id = ?", (new_id,))
            db.executemany("""INSERT INTO prod_parts (org_id, order_id, name, company, months, serial, sort)
                              VALUES (?,?,?,?,?,?,?)""", [(oid, new_id, p["name"], p["company"], p["months"],
                                                           p["serial"], i) for i, p in enumerate(l["parts"])])
        if l["bill_no"]:
            bills.add(l["bill_no"])
        done += 1
    db.commit()
    try:
        os.remove(path)
    except OSError:
        pass
    H.audit("Production: old deliveries imported", f"{done} orders")
    flash(f"{done} old deliver{'ies' if done != 1 else 'y'} imported (Delivered, warranty from the bill date).",
          "success")
    return redirect(url_for("prod.orders", show="delivered"))
