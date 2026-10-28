"""Cafe module: org type, permissions, menu, billing maths, KOT, online orders, points, referral, expenses, analysis,
tenant isolation, and that existing CRM organizations are untouched.

Run:  python -m pytest -q tests/test_cafe.py
"""
import io
import json
import re
import sqlite3

import pytest

from test_crm import PW, Browser, crm, crm_db  # noqa: F401  (reuses the CRM test helpers)

import cafe


@pytest.fixture()
def cw(tmp_path, monkeypatch):
    db_path = str(tmp_path / "crm.db")
    monkeypatch.setattr(crm, "DB_PATH", db_path)
    crm._failed_logins.clear()
    import cafe_pos
    cafe_pos._order_hits.clear()
    crm_db.migrate(db_path)
    sup = Browser()
    sup.post("/setup", {"agency": "Real View", "name": "Owner", "username": "super", "password": PW})

    def make_org(name, code, admin_user, app_type):
        r = sup.post("/admin/organizations/new", {
            "name": name, "code": code, "admin_name": f"{name} Admin", "admin_username": admin_user,
            "onboarding": "password", "password": "TempPass123", "app_type": app_type})
        assert f"created with admin {admin_user}" in r.text
        b = Browser()
        b.login(admin_user, "TempPass123")
        b.post("/account", {"current_password": "TempPass123", "new_password": PW, "confirm": PW})
        return b

    W = dict(sup=sup, db_path=db_path)
    W["cafe"] = make_org("Chai Point", "CHAI", "chai_admin", "cafe")
    W["cafe2"] = make_org("Other Cafe", "OTHER", "other_admin", "cafe")
    W["crm"] = make_org("Machines Ltd", "MACH", "mach_admin", "crm")
    W["both"] = make_org("Both Co", "BOTH", "both_admin", "both")

    def db():
        c = sqlite3.connect(db_path)
        c.row_factory = sqlite3.Row
        return c
    W["db"] = db
    W["oid"] = lambda code: db().execute("SELECT id FROM organizations WHERE code = ?", (code,)).fetchone()[0]
    return W


def make_staff(admin, username, role="cashier", perms=None):
    data = {"name": username.title(), "username": username, "role": "staff", "active": "1", "onboarding": "password",
            "password": "TempPass123", "cafe_role": role}
    if perms is not None:
        data["cafe_perm"] = perms
    r = admin.post("/staff/new", data)
    assert "created" in r.text, r.text[:2000]
    b = Browser()
    b.login(username, "TempPass123")
    b.post("/account", {"current_password": "TempPass123", "new_password": PW, "confirm": PW})
    return b


def png(size=(1600, 1200), color=(200, 30, 30)):
    from PIL import Image
    buf = io.BytesIO()
    import random
    img = Image.new("RGB", size, color)
    px = img.load()
    for _ in range(40000):  # noise so the JPEG is not tiny
        px[random.randrange(size[0]), random.randrange(size[1])] = tuple(random.randrange(256) for _ in range(3))
    img.save(buf, "PNG")
    buf.seek(0)
    return buf


def add_item(b, name, price, **extra):
    data = {"name": name, "price": price, "available": "1"}
    data.update(extra)
    r = b.post("/cafe/menu/item/new", data)
    assert f"&#39;{name}&#39; added to the menu" in r.text, re.findall(r'flash-\w+">([^<]*)', r.text)


def item_id(w, name, code="CHAI"):
    return w["db"]().execute("SELECT id FROM menu_items WHERE name = ? AND org_id = ?", (name, w["oid"](code))).fetchone()[0]


def pos(b, payload):
    r = b.post("/cafe/pos/save", {"payload": json.dumps(payload)}, follow=False)
    return r.status_code, r.get_json()


def settings(b, **kw):
    """Save cafe settings with sensible defaults (checkboxes need to be sent every time)."""
    data = {"name": "Chai Point", "paper": "80", "kot_on": "1", "kot_with_bill": "1", "menu_on": "1",
            "gst_inclusive": "1", "default_gst": "5", "order_takeaway": "1", "order_dine_in": "1",
            "loyalty_mode": "percent", "loyalty_percent": "5", "loyalty_point_value": "1", "loyalty_max_redeem_pct": "50",
            "sort_pos": "popular", "sort_menu": "popular"}
    data.update(kw)
    data = {k: v for k, v in data.items() if v is not None}
    r = b.post("/cafe/settings", data)
    assert "Settings saved" in r.text, r.text[:1500]


# ------------------------------------------------------------------ org types & navigation

def test_org_types_and_existing_orgs_stay_crm(cw):
    w = cw
    rows = {r["code"]: r["app_type"] for r in w["db"]().execute("SELECT code, app_type FROM organizations")}
    assert rows == {"AGENCY": "crm", "CHAI": "cafe", "OTHER": "cafe", "MACH": "crm", "BOTH": "both"}
    # cafe-only org: CRM pages lead to the cafe; dashboard shows the cafe
    r = w["cafe"].get("/", follow_redirects=True)
    assert "New Bill" in r.text and "Leads" not in r.text.split("<main")[0]
    assert w["cafe"].get("/leads").headers["Location"].endswith("/cafe")
    # CRM org cannot open cafe pages
    assert w["crm"].get("/cafe").status_code == 404
    assert w["crm"].get("/cafe/pos").status_code == 404
    assert "Leads" in w["crm"].get("/").text
    # both: switch
    r = w["both"].get("/")
    assert "mode-switch" in r.text and "Leads" in r.text
    r = w["both"].post("/switch/cafe")
    assert "New Bill" in r.text
    r = w["both"].get("/", follow_redirects=True)
    assert "New Bill" in r.text
    w["both"].post("/switch/crm")
    assert "Today's Leads" in w["both"].get("/").text
    # Super Admin changes the type of an existing org without losing anything
    oid = w["oid"]("MACH")
    r = w["sup"].post(f"/admin/organizations/{oid}/edit", {"name": "Machines Ltd", "code": "MACH", "app_type": "both",
                                                          "wa_templates_allowed": "1", "catalog_allowed": "1"})
    assert "Organization updated" in r.text
    assert w["db"]().execute("SELECT app_type FROM organizations WHERE id = ?", (oid,)).fetchone()[0] == "both"


# ------------------------------------------------------------------ menu & photos

def test_menu_items_photos_under_50kb_and_offers(cw):
    w = cw
    b = w["cafe"]
    b.post("/cafe/menu/category", {"action": "add", "name": "Juices"})
    cat = w["db"]().execute("SELECT id FROM menu_categories WHERE name = 'Juices'").fetchone()[0]
    r = b.post("/cafe/menu/item/new", {"name": "Mango Juice", "price": "80", "offer_price": "60", "category_id": cat,
                                       "available": "1", "food_type": "veg", "tags": ["bestseller"],
                                       "photo": (png(), "big.png")}, content_type="multipart/form-data")
    assert "&#39;Mango Juice&#39; added to the menu" in r.text
    it = w["db"]().execute("SELECT * FROM menu_items WHERE name = 'Mango Juice'").fetchone()
    assert it["price"] == 8000 and it["offer_price"] == 6000 and it["photo"]
    import os
    size = os.path.getsize(os.path.join(crm.media_dir(), it["photo"]))
    assert size <= 50 * 1024, size
    from PIL import Image
    assert Image.open(os.path.join(crm.media_dir(), it["photo"])).size == (600, 600)
    # hint is shown on the form
    assert "600 × 600" in b.get("/cafe/menu/item/new").text
    # offer price must be lower than the price
    r = b.post("/cafe/menu/item/new", {"name": "Bad", "price": "50", "offer_price": "60", "available": "1"})
    assert "must be lower than the price" in r.text
    # without an offer price no offer is shown
    add_item(b, "Plain Tea", "15")
    settings(b)
    page = w["sup"].get("/m/chai")  # public, no login needed
    assert page.status_code == 200
    data = json.loads(re.search(r'id="menu-cfg" type="application/json">(.*?)</script>', page.text, re.S).group(1))
    items = {i["name"]: i for i in data["items"]}
    assert items["Mango Juice"]["price"] == 6000 and items["Mango Juice"]["mrp"] == 8000 and items["Mango Juice"]["off"] == 25
    assert items["Plain Tea"]["mrp"] is None
    # offer items become banners automatically when no banner exists
    assert data["banners"] and data["banners"][0]["auto"]
    # an expired offer goes back to the normal price
    iid = item_id(w, "Mango Juice")
    w["db"]().execute("UPDATE menu_items SET offer_to = '2020-01-01' WHERE id = ?", (iid,)).connection.commit()
    data = json.loads(re.search(r'id="menu-cfg" type="application/json">(.*?)</script>',
                                w["sup"].get("/m/chai").text, re.S).group(1))
    assert {i["name"]: i for i in data["items"]}["Mango Juice"]["price"] == 8000
    # banner with a wide photo
    r = b.post("/cafe/banners/save", {"title": "Weekend ₹99", "active": "1", "photo": (png((2000, 800)), "b.png")},
               content_type="multipart/form-data")
    assert "Banner added" in r.text and "1200 × 500" in r.text
    ph = w["db"]().execute("SELECT photo FROM menu_banners").fetchone()[0]
    assert os.path.getsize(os.path.join(crm.media_dir(), ph)) <= 50 * 1024
    assert Image.open(os.path.join(crm.media_dir(), ph)).size == (1200, 500)
    # QR code
    assert w["cafe"].get("/cafe/qr.svg").data.startswith(b"<svg")
    assert w["cafe"].get("/cafe/qr.png").data[:4] == b"\x89PNG"
    # menu off -> 404 for the public
    settings(b, menu_on=None)
    assert w["sup"].get("/m/chai").status_code == 404
    assert b.get("/m/chai?preview=1").status_code == 200  # the cafe's own users can preview


# ------------------------------------------------------------------ billing maths

def test_compute_totals():
    lines = [{"price": 10000, "qty": 2, "gst_rate": 5}, {"price": 5000, "qty": 1, "gst_rate": 18}]
    t = cafe.compute_totals(lines, False, True, 0)
    assert t["subtotal"] == 25000 and t["gross"] == 25000 and t["cgst"] == 0
    t = cafe.compute_totals(lines, True, False, 0)  # exclusive: tax added
    assert t["taxable"] == 25000 and t["cgst"] + t["sgst"] == 1000 + 900 and t["gross"] == 26900
    t = cafe.compute_totals(lines, True, True, 0)  # inclusive: total unchanged
    assert t["gross"] == 25000 and t["taxable"] + t["cgst"] + t["sgst"] == 25000
    t = cafe.compute_totals(lines, True, False, 5000)  # discount before tax
    assert t["discount"] == 5000 and t["taxable"] == 20000
    assert cafe.rupees(123456789) == "₹12,34,567.89" and cafe.rupees(10000) == "₹100"
    assert cafe.to_paise("120.5") == 12050 and cafe.to_paise("₹1,200") == 120000 and cafe.to_paise("abc") is None


def test_takeaway_bill_gst_discount_kot_and_print(cw):
    w = cw
    b = w["cafe"]
    settings(b, gst_on="1", gstin="32ABCDE1234F1Z5", gst_inclusive=None, bill_prefix="B-")
    add_item(b, "Shawarma", "100", gst_rate="5")
    add_item(b, "Lime", "30", gst_rate="5")
    sh, li = item_id(w, "Shawarma"), item_id(w, "Lime")
    code, d = pos(b, {"action": "bill", "order_type": "takeaway", "payment_mode": "UPI",
                      "items": [{"item_id": sh, "qty": 2}, {"item_id": li, "qty": 1}],
                      "discount": {"mode": "amt", "value": "30"}})
    assert code == 200 and d["ok"], d
    bill = w["db"]().execute("SELECT * FROM cafe_bills").fetchone()
    # 230 - 30 = 200 taxable, 5% = 10 -> 210
    assert (bill["subtotal"], bill["discount"], bill["taxable"], bill["cgst"], bill["sgst"], bill["total"]) == \
        (23000, 3000, 20000, 500, 500, 21000)
    assert bill["bill_no"] == 1 and bill["payment_mode"] == "UPI" and bill["status"] == "paid"
    # KOT printed together with the takeaway bill
    assert "kot=" in d["print"]
    r = b.get(d["print"])
    assert "TAX INVOICE" in r.text and "B-1" in r.text and "GSTIN: 32ABCDE1234F1Z5" in r.text and "KOT #1" in r.text
    # paper sizes
    assert "size: 58mm" in b.get(d["print"].replace("auto=1", "paper=58")).text
    # public bill link and WhatsApp text
    view = b.get(d["view"]).text
    assert "wa.me" in view
    tok = bill["token"]
    pub = Browser().get(f"/b/{tok}")
    assert pub.status_code == 200 and "₹210" in pub.text
    # payment mode is required, items required
    code, d2 = pos(b, {"action": "bill", "order_type": "takeaway", "items": [{"item_id": sh, "qty": 1}]})
    assert code == 400
    code, d2 = pos(b, {"action": "bill", "order_type": "takeaway", "payment_mode": "Cash", "items": []})
    assert code == 400
    # next bill number
    pos(b, {"action": "bill", "order_type": "takeaway", "payment_mode": "Cash", "items": [{"item_id": li, "qty": 1}]})
    assert [r[0] for r in w["db"]().execute("SELECT bill_no FROM cafe_bills ORDER BY id")] == [1, 2]


def test_dine_in_table_kot_flow_and_cancel(cw):
    w = cw
    b = w["cafe"]
    settings(b)
    add_item(b, "Tea", "15")
    add_item(b, "Puffs", "25")
    tea, puff = item_id(w, "Tea"), item_id(w, "Puffs")
    b.post("/cafe/tables/save", {"action": "bulk", "count": "3", "seats": "4"})
    t1 = w["db"]().execute("SELECT id FROM cafe_tables WHERE name = 'T1'").fetchone()[0]
    code, d = pos(b, {"action": "kot", "order_type": "dine_in", "table_id": t1, "items": [{"item_id": tea, "qty": 2}]})
    assert d["ok"] and "/cafe/kot/" in d["print"]
    oid_ = d["order_id"]
    assert "T1" in b.get("/cafe/tables").text and "busy" in b.get("/cafe/tables").text
    # table is taken: a second new order on it is refused
    code, d2 = pos(b, {"action": "kot", "order_type": "dine_in", "table_id": t1, "items": [{"item_id": tea, "qty": 1}]})
    assert code == 400
    # add more items -> the next KOT has only the new ones
    code, d = pos(b, {"action": "kot", "order_id": oid_, "order_type": "dine_in", "table_id": t1,
                      "items": [{"item_id": tea, "qty": 3}, {"item_id": puff, "qty": 1}]})
    kots = w["db"]().execute("SELECT lines FROM cafe_kots ORDER BY id").fetchall()
    assert json.loads(kots[1][0]) == [{"name": "Tea", "qty": 1, "note": ""}, {"name": "Puffs", "qty": 1, "note": ""}]
    # nothing new -> no KOT
    code, d = pos(b, {"action": "kot", "order_id": oid_, "order_type": "dine_in", "table_id": t1,
                      "items": [{"item_id": tea, "qty": 3}, {"item_id": puff, "qty": 1}]})
    assert code == 400
    # bill the table
    code, d = pos(b, {"action": "bill", "order_id": oid_, "order_type": "dine_in", "table_id": t1, "payment_mode": "Cash",
                      "items": [{"item_id": tea, "qty": 3}, {"item_id": puff, "qty": 1}]})
    assert d["ok"]
    bill = w["db"]().execute("SELECT * FROM cafe_bills").fetchone()
    assert bill["total"] == 7000 and bill["table_name"] == "T1"
    assert "free" in b.get("/cafe/tables").text
    # cashier cannot cancel; admin can, with a reason
    cashier = make_staff(b, "cash1")
    assert cashier.post(f"/cafe/bills/{bill['id']}/cancel", {"reason": "x"}, follow=False).status_code == 403
    r = b.post(f"/cafe/bills/{bill['id']}/cancel", {"reason": ""})
    assert "reason for cancelling" in r.text
    r = b.post(f"/cafe/bills/{bill['id']}/cancel", {"reason": "Wrong table"})
    assert "cancelled. Points" in r.text
    assert w["db"]().execute("SELECT status FROM cafe_bills").fetchone()[0] == "cancelled"


# ------------------------------------------------------------------ permissions

def test_staff_permissions_and_hidden_menus(cw):
    w = cw
    b = w["cafe"]
    cashier = make_staff(b, "cashier1", "cashier")
    page = cashier.get("/cafe").text
    assert "New Bill" in page and "Analysis" not in page and "Settings" not in page and "Expenses" not in page
    for url in ("/cafe/analysis", "/cafe/settings", "/cafe/menu", "/cafe/expenses", "/cafe/templates"):
        assert cashier.get(url).status_code == 403, url
    assert cashier.get("/cafe/pos").status_code == 200
    # custom: only expenses
    exp = make_staff(b, "expenser", "custom", ["expense_add"])
    assert exp.get("/cafe/expenses").status_code == 200
    assert exp.get("/cafe/pos").status_code == 403
    r = exp.post("/cafe/expenses/save", {"kind": "expense", "amount": "500", "category": "Gas", "payment_mode": "Cash"})
    assert "₹500" in r.text
    assert exp.get("/cafe/analysis").status_code == 403
    # a staff manager can add staff but cannot hand out permissions they don't have, nor create admins
    mgr = make_staff(b, "mgr", "custom", ["staff_manage", "bill_create"])
    assert mgr.get("/staff").status_code == 200
    r = mgr.post("/staff/new", {"name": "Sneaky", "username": "sneaky", "role": "org_admin", "active": "1",
                                "onboarding": "password", "password": "TempPass123", "cafe_role": "custom",
                                "cafe_perm": ["settings_edit", "bill_create"]})
    row = w["db"]().execute("SELECT role, cafe_perms FROM users WHERE username = 'sneaky'").fetchone()
    assert row["role"] == "staff" and json.loads(row["cafe_perms"]) == ["bill_create"]
    admin_id = w["db"]().execute("SELECT id FROM users WHERE username = 'chai_admin'").fetchone()[0]
    assert mgr.get(f"/staff/{admin_id}/edit").status_code == 403
    # role presets fill the ticks
    perms = json.loads(w["db"]().execute("SELECT cafe_perms FROM users WHERE username = 'cashier1'").fetchone()[0])
    assert set(perms) == set(cafe.ROLE_PRESETS["cashier"])
    # discount needs its permission
    add_item(b, "Tea", "15")
    code, d = pos(cashier, {"action": "bill", "order_type": "takeaway", "payment_mode": "Cash",
                            "items": [{"item_id": item_id(w, "Tea"), "qty": 1}], "discount": {"mode": "amt", "value": "5"}})
    assert code == 400 and "discount" in d["error"]


# ------------------------------------------------------------------ online orders

def test_online_order_pending_confirm_and_status(cw):
    w = cw
    b = w["cafe"]
    settings(b, online_on="1")
    add_item(b, "Tea", "15")
    add_item(b, "Cake", "60", available=None)  # switched off
    b.post("/cafe/tables/save", {"name": "T9"})
    tea, cake = item_id(w, "Tea"), item_id(w, "Cake")
    t9 = w["db"]().execute("SELECT id FROM cafe_tables WHERE name = 'T9'").fetchone()[0]
    guest = Browser()
    menu = guest.get(f"/m/chai?t={t9}")
    assert menu.status_code == 200 and "You're at <b>T9</b>" in menu.text

    def order(payload):
        r = guest.post("/m/chai/order", {"payload": json.dumps(payload)}, follow=False)
        return r.status_code, r.get_json()

    code, d = order({"type": "dine_in", "table": t9, "name": "Anu", "phone": "9876543210", "items": [{"item_id": cake, "qty": 1}]})
    assert code == 400 and "not available" in d["error"]
    code, d = order({"type": "dine_in", "table": t9, "name": "Anu", "phone": "123", "items": [{"item_id": tea, "qty": 1}]})
    assert code == 400
    code, d = order({"type": "dine_in", "table": t9, "name": "Anu", "phone": "+91 98765 43210", "items": [{"item_id": tea, "qty": 2}]})
    assert code == 200 and d["ok"]
    status = guest.get(d["url"])
    assert status.status_code == 200 and "Waiting for the cafe" in status.text
    assert guest.get(d["url"] + ".json").get_json()["status"] == "pending"
    # pending: not a bill, not counted
    assert b.get("/cafe/orders/pending.json").get_json()["count"] == 1
    assert w["db"]().execute("SELECT COUNT(*) FROM cafe_bills").fetchone()[0] == 0
    assert "Anu" in b.get("/cafe/orders").text
    oid_ = w["db"]().execute("SELECT id FROM cafe_orders WHERE status = 'pending'").fetchone()[0]
    r = b.post(f"/cafe/orders/{oid_}/confirm")
    assert "confirmed" in r.text
    assert guest.get(d["url"] + ".json").get_json()["status"] == "confirmed"
    o = w["db"]().execute("SELECT * FROM cafe_orders WHERE id = ?", (oid_,)).fetchone()
    assert o["status"] == "open" and o["customer_id"]
    assert w["db"]().execute("SELECT COUNT(*) FROM cafe_kots WHERE order_id = ?", (oid_,)).fetchone()[0] == 1
    b.post(f"/cafe/orders/{oid_}/ready")
    assert guest.get(d["url"] + ".json").get_json()["status"] == "ready"
    code, bd = pos(b, {"action": "bill", "order_id": oid_, "order_type": "dine_in", "table_id": t9, "payment_mode": "Cash",
                       "items": [{"item_id": tea, "qty": 2}]})
    assert bd["ok"]
    assert guest.get(d["url"] + ".json").get_json()["status"] == "done"
    # reject
    code, d = order({"type": "takeaway", "name": "Ben", "phone": "9000000001", "items": [{"item_id": tea, "qty": 1}]})
    oid2 = w["db"]().execute("SELECT id FROM cafe_orders WHERE status = 'pending'").fetchone()[0]
    b.post(f"/cafe/orders/{oid2}/reject", {"reason": "Closing"})
    s = guest.get(d["url"] + ".json").get_json()
    assert s == {"status": "rejected", "reason": "Closing"}
    # online ordering off -> refused
    settings(b, online_on=None)
    code, d = order({"type": "takeaway", "name": "Ben", "phone": "9000000001", "items": [{"item_id": tea, "qty": 1}]})
    assert code == 400
    # cashier without the online-orders tick cannot confirm
    settings(b, online_on="1")
    clerk = make_staff(b, "clerk", "custom", ["bill_create"])
    code, d = order({"type": "takeaway", "name": "Cy", "phone": "9000000002", "items": [{"item_id": tea, "qty": 1}]})
    oid3 = w["db"]().execute("SELECT id FROM cafe_orders WHERE status = 'pending'").fetchone()[0]
    assert clerk.post(f"/cafe/orders/{oid3}/confirm", follow=False).status_code == 403


# ------------------------------------------------------------------ customers, points, referral

def test_points_earn_redeem_cancel_referral_and_expiry(cw):
    w = cw
    b = w["cafe"]
    settings(b, loyalty_on="1", loyalty_percent="10", loyalty_point_value="0.50", loyalty_min_redeem="10",
             loyalty_max_redeem_pct="20", referral_on="1", referral_referrer_points="50", referral_new_points="20",
             birthday_on="1", loyalty_expiry_days="30")
    add_item(b, "Meal", "200")
    meal = item_id(w, "Meal")
    # first customer: ₹400 -> 40 points
    code, d = pos(b, {"action": "bill", "order_type": "takeaway", "payment_mode": "Cash", "items": [{"item_id": meal, "qty": 2}],
                      "customer": {"phone": "9847000001", "name": "Riya", "birthday": "1995-10-24"}})
    assert d["ok"], d
    db = w["db"]()
    riya = db.execute("SELECT * FROM customers WHERE wa_norm = '9847000001'").fetchone()
    assert riya["points"] == 40 and riya["visits"] == 1 and riya["total_spent"] == 40000 and riya["birthday"] == "10-24"
    # lookup for billing
    lk = b.get("/cafe/customers/lookup.json?phone=9847000001&payable=20000").get_json()
    assert lk["found"] and lk["points"] == 40 and lk["max_points"] == 40  # 20% of ₹200 = ₹40 = 80 pts, capped by balance
    # redeem more than allowed -> refused; 30 points = ₹15 off ₹200 -> ₹185, earns 18
    code, d = pos(b, {"action": "bill", "order_type": "takeaway", "payment_mode": "Cash", "items": [{"item_id": meal, "qty": 1}],
                      "customer": {"phone": "9847000001"}, "redeem_points": 41})
    assert code == 400
    code, d = pos(b, {"action": "bill", "order_type": "takeaway", "payment_mode": "Cash", "items": [{"item_id": meal, "qty": 1}],
                      "customer": {"phone": "9847000001"}, "redeem_points": 30})
    assert d["ok"], d
    bill = db.execute("SELECT * FROM cafe_bills ORDER BY id DESC").fetchone()
    assert bill["points_value"] == 1500 and bill["total"] == 18500 and bill["points_earned"] == 18
    assert db.execute("SELECT points FROM customers WHERE id = ?", (riya["id"],)).fetchone()[0] == 40 - 30 + 18
    # cancel restores
    b.post(f"/cafe/bills/{bill['id']}/cancel", {"reason": "test"})
    assert db.execute("SELECT points FROM customers WHERE id = ?", (riya["id"],)).fetchone()[0] == 40
    # referral: new customer with Riya's code; points only after the first bill
    b.post("/cafe/customers/new", {"name": "Sam", "whatsapp": "9847000002", "ref": riya["referral_code"]})
    sam = db.execute("SELECT * FROM customers WHERE wa_norm = '9847000002'").fetchone()
    assert sam["referred_by"] == riya["id"] and sam["points"] == 0
    pos(b, {"action": "bill", "order_type": "takeaway", "payment_mode": "UPI", "items": [{"item_id": meal, "qty": 1}],
            "customer": {"phone": "9847000002"}})
    assert db.execute("SELECT points FROM customers WHERE id = ?", (riya["id"],)).fetchone()[0] == 40 + 50
    assert db.execute("SELECT points FROM customers WHERE id = ?", (sam["id"],)).fetchone()[0] == 20 + 20
    pos(b, {"action": "bill", "order_type": "takeaway", "payment_mode": "UPI", "items": [{"item_id": meal, "qty": 1}],
            "customer": {"phone": "9847000002"}})
    assert db.execute("SELECT points FROM customers WHERE id = ?", (riya["id"],)).fetchone()[0] == 90  # only once
    # manual adjust needs a reason; cashier cannot
    r = b.post(f"/cafe/customers/{riya['id']}/points", {"points": "10", "direction": "add", "reason": ""})
    assert "Points and a reason are required" in r.text
    b.post(f"/cafe/customers/{riya['id']}/points", {"points": "10", "direction": "sub", "reason": "Gift"})
    assert db.execute("SELECT points FROM customers WHERE id = ?", (riya["id"],)).fetchone()[0] == 80
    cashier = make_staff(b, "cashier2")
    assert cashier.post(f"/cafe/customers/{riya['id']}/points", {"points": "10", "direction": "add", "reason": "x"},
                        follow=False).status_code == 403
    # expiry: make all of Riya's points old
    db.execute("UPDATE points_ledger SET expires_at = '2020-01-01' WHERE customer_id = ? AND remaining > 0", (riya["id"],))
    db.commit()
    b.get("/cafe/customers")
    assert db.execute("SELECT points FROM customers WHERE id = ?", (riya["id"],)).fetchone()[0] == 0
    # marketing filters and template links
    page = b.get("/cafe/customers?f=birthday").text
    tpl = db.execute("SELECT id FROM cafe_templates WHERE title = 'Offer'").fetchone()[0]
    page = b.get(f"/cafe/customers?tpl={tpl}").text
    assert "wa.me/919847000002?text=" in page and "Sam" in page


# ------------------------------------------------------------------ expenses & analysis

def test_expenses_and_analysis(cw):
    w = cw
    b = w["cafe"]
    settings(b)
    add_item(b, "Tea", "20")
    tea = item_id(w, "Tea")
    pos(b, {"action": "bill", "order_type": "takeaway", "payment_mode": "Cash", "items": [{"item_id": tea, "qty": 10}]})
    b.post("/cafe/expenses/save", {"kind": "expense", "amount": "50", "category": "Gas", "payment_mode": "Cash"})
    b.post("/cafe/expenses/save", {"kind": "expense", "amount": "30.50", "category": "Raw materials", "payment_mode": "UPI"})
    b.post("/cafe/expenses/save", {"kind": "income", "amount": "100", "category": "Catering", "payment_mode": "Cash"})
    r = b.post("/cafe/expenses/save", {"kind": "expense", "amount": "-5", "category": "Gas"})
    assert "valid amount" in r.text
    page = b.get("/cafe/analysis?p=today").text
    assert "₹300" in page  # income 200 + 100
    assert "₹80" in page  # expense 80.50 shown rounded
    assert "<svg" in page and "data-tip" in page
    csv = b.get("/cafe/export/daily.csv?p=today").data.decode("utf-8-sig")
    assert "200.00" in csv and "80.50" in csv and "219.50" in csv
    assert b.get("/cafe/export/bills.csv?p=today").status_code == 200
    assert b.get("/cafe/export/customers.csv").status_code == 200


# ------------------------------------------------------------------ isolation

def test_cafe_tenant_isolation(cw):
    w = cw
    a, other = w["cafe"], w["cafe2"]
    settings(a)
    add_item(a, "Secret Tea", "15")
    tea = item_id(w, "Secret Tea")
    code, d = pos(a, {"action": "bill", "order_type": "takeaway", "payment_mode": "Cash", "items": [{"item_id": tea, "qty": 1}],
                      "customer": {"phone": "9000011111", "name": "Private Person"}})
    bill_id = w["db"]().execute("SELECT id FROM cafe_bills").fetchone()[0]
    cust_id = w["db"]().execute("SELECT id FROM customers").fetchone()[0]
    assert other.get(f"/cafe/bills/{bill_id}").status_code == 404
    assert other.get(f"/cafe/customers/{cust_id}").status_code == 404
    assert other.get(f"/cafe/menu/item/{tea}/edit").status_code == 404
    assert "Private Person" not in other.get("/cafe/customers").text
    # using another cafe's item id in a bill is refused
    code, d = pos(other, {"action": "bill", "order_type": "takeaway", "payment_mode": "Cash", "items": [{"item_id": tea, "qty": 1}]})
    assert code == 400
    assert not other.get("/cafe/customers/lookup.json?phone=9000011111").get_json()["found"]
    # database guard
    db = w["db"]()
    with pytest.raises(sqlite3.IntegrityError):
        db.execute("INSERT INTO cafe_order_items (org_id, order_id, item_id, name, price, qty) VALUES (?,?,?,?,?,?)",
                   (w["oid"]("OTHER"), d and 1 or 1, tea, "x", 1, 1))
    # Super Admin viewing a cafe is read-only
    w["sup"].post(f"/admin/organizations/{w['oid']('CHAI')}/view")
    assert "Secret Tea" in w["sup"].get("/cafe/menu").text
    assert w["sup"].post("/cafe/menu/category", {"action": "add", "name": "X"}, follow=False).status_code == 403


def test_migration_from_v7_keeps_crm_data(tmp_path, monkeypatch):
    """An existing v7 database (with leads) upgrades to v8 with a backup and without losing anything."""
    db_path = str(tmp_path / "old.db")
    monkeypatch.setattr(crm, "DB_PATH", db_path)
    saved = list(crm_db.MIGRATIONS)
    monkeypatch.setattr(crm_db, "MIGRATIONS", [m for m in saved if m[0] <= 7])
    crm_db.migrate(db_path)
    sup = Browser()
    crm._failed_logins.clear()
    sup.post("/setup", {"agency": "Old Agency", "name": "Owner", "username": "super", "password": PW})
    conn = sqlite3.connect(db_path)
    oid = conn.execute("SELECT id FROM organizations").fetchone()[0]
    uid = conn.execute("SELECT id FROM users").fetchone()[0]
    pid = conn.execute("INSERT INTO projects (org_id, name, status, created_by, created_at) VALUES (?, 'P', 'Active', ?, 'x')",
                       (oid, uid)).lastrowid
    conn.execute("INSERT INTO leads (org_id, project_id, customer_name, entered_by, created_at, updated_at)"
                 " VALUES (?,?, 'Old Lead', ?, 'x', 'x')", (oid, pid, uid))
    conn.commit()
    assert conn.execute("PRAGMA user_version").fetchone()[0] == 7
    conn.close()
    monkeypatch.setattr(crm_db, "MIGRATIONS", saved)
    backup = crm_db.migrate(db_path)
    assert backup and "pre-v8" in backup
    conn = sqlite3.connect(db_path)
    assert conn.execute("PRAGMA user_version").fetchone()[0] == crm_db.SCHEMA_VERSION
    assert conn.execute("SELECT customer_name FROM leads").fetchone()[0] == "Old Lead"
    assert conn.execute("SELECT app_type FROM organizations").fetchone()[0] == "crm"
    assert sqlite3.connect(backup).execute("SELECT COUNT(*) FROM leads").fetchone()[0] == 1


# ------------------------------------------------------------------ serial numbers and Excel upload

def xlsx(rows):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Menu"
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_serial_numbers(cw):
    w = cw
    b = w["cafe"]
    settings(b)
    add_item(b, "Tea", "15")
    add_item(b, "Coffee", "20")
    rows = {r["name"]: r["serial"] for r in w["db"]().execute("SELECT name, serial FROM menu_items")}
    assert rows == {"Tea": 1, "Coffee": 2}
    r = b.post("/cafe/menu/item/new", {"name": "Juice", "price": "40", "serial": "1", "available": "1"})
    assert "S.No 1 is already used by" in r.text
    add_item(b, "Juice", "40", serial="101")
    assert w["db"]().execute("SELECT serial FROM menu_items WHERE name = 'Juice'").fetchone()[0] == 101
    # staff see S.No in billing; customers do not
    assert '"sno": 101' in b.get("/cafe/pos").text
    assert '"sno"' not in Browser().get("/m/chai").text
    assert "Search item or S.No" in b.get("/cafe/menu").text
    page = b.get("/cafe/menu?q=101").text
    assert ">Juice<" in page and ">Coffee<" not in page


def test_excel_template_export_import_and_photos(cw):
    w = cw
    b = w["cafe"]
    settings(b)
    add_item(b, "Tea", "15")
    from openpyxl import load_workbook
    t = b.get("/cafe/menu/template.xlsx")
    assert t.status_code == 200 and "spreadsheetml" in t.headers["Content-Type"]
    ws = load_workbook(io.BytesIO(t.data))["Menu"]
    assert [c.value for c in ws[1]][:4] == ["S.No", "Category", "Item name", "Price"]
    e = load_workbook(io.BytesIO(b.get("/cafe/menu/export.xlsx").data))["Menu"]
    assert e.cell(2, 3).value == "Tea" and e.cell(2, 4).value == 15 and e.cell(2, 1).value == 1
    sheet = xlsx([["S.No", "Category", "Item name", "Price", "Offer price", "Veg / Non-veg", "GST %", "Description", "Available"],
                  [1, "Hot drinks", "Tea", 18, None, "Veg", 5, "Masala tea", "Yes"],          # update by S.No
                  [None, "Juices", "Mango Shake", 80, 65, "Veg", 5, "", "Yes"],               # new, auto S.No
                  [50, "Snacks", "Chicken Puffs", "25", None, "Non-veg", None, "", "No"],     # new with S.No 50
                  [None, "Snacks", "", 30, None, "", None, "", ""],                           # error: no name
                  [None, "Snacks", "Bad Price", "abc", None, "", None, "", ""],               # error: price
                  [None, "Snacks", "Cheap", 10, 20, "", None, "", ""]])                       # error: offer >= price
    r = b.post("/cafe/menu/import", {"sheet": (sheet, "menu.xlsx")}, content_type="multipart/form-data")
    assert "Import preview" in r.text and "item name is empty" in r.text and "price is missing" in r.text
    assert "offer price must be lower" in r.text
    assert w["db"]().execute("SELECT COUNT(*) FROM menu_items").fetchone()[0] == 1  # nothing saved yet
    tok = re.search(r'name="tok" value="([0-9a-f]+)"', r.text).group(1)
    # another cafe cannot use this preview
    assert w["cafe2"].post("/cafe/menu/import/confirm", {"tok": tok}, follow=False).status_code == 404
    r = b.post("/cafe/menu/import/confirm", {"tok": tok})
    assert "2 new items added, 1 items updated" in r.text
    db = w["db"]()
    items = {r["name"]: r for r in db.execute("SELECT * FROM menu_items WHERE active = 1")}
    assert items["Tea"]["price"] == 1800 and items["Tea"]["description"] == "Masala tea" and items["Tea"]["serial"] == 1
    assert items["Mango Shake"]["offer_price"] == 6500 and items["Mango Shake"]["serial"] == 2
    assert items["Chicken Puffs"]["serial"] == 50 and items["Chicken Puffs"]["available"] == 0
    assert items["Chicken Puffs"]["food_type"] == "nonveg"
    cats = {r[0] for r in db.execute("SELECT name FROM menu_categories")}
    assert {"Hot drinks", "Juices", "Snacks"} <= cats
    # the preview can be used once only
    r = b.post("/cafe/menu/import/confirm", {"tok": tok})
    assert "expired" in r.text
    # wrong file type
    r = b.post("/cafe/menu/import", {"sheet": (io.BytesIO(b"hello"), "menu.pdf")}, content_type="multipart/form-data")
    assert "Upload the Excel file" in r.text
    # photos matched by item name and by S.No
    r = b.post("/cafe/menu/photos", {"photos": [(png(), "Mango Shake.JPG"), (png(), "50.png"), (png(), "Unknown.jpg")]},
               content_type="multipart/form-data")
    assert "2 photos added" in r.text and "Unknown.jpg" in r.text
    import os
    for name in ("Mango Shake", "Chicken Puffs"):
        ph = db.execute("SELECT photo FROM menu_items WHERE name = ?", (name,)).fetchone()[0]
        assert ph and os.path.getsize(os.path.join(crm.media_dir(), ph)) <= 50 * 1024
    # cashier cannot import
    cashier = make_staff(b, "cash9")
    assert cashier.get("/cafe/menu/bulk").status_code == 403


# ------------------------------------------------------------------ round 3: units, price change, barcode, order, BT

def test_weighed_items_kg_litre_bill_kot_and_reports(cw):
    w = cw
    b = w["cafe"]
    settings(b)
    add_item(b, "Mixture", "480", unit="kg")
    add_item(b, "Milk", "60", unit="l")
    add_item(b, "Puffs", "20")
    mx, mk, pf = item_id(w, "Mixture"), item_id(w, "Milk"), item_id(w, "Puffs")
    assert w["db"]().execute("SELECT unit FROM menu_items WHERE id = ?", (mx,)).fetchone()[0] == "kg"
    code, d = pos(b, {"action": "bill", "order_type": "takeaway", "payment_mode": "Cash",
                      "items": [{"item_id": mx, "qty": 1, "measure": 0.25}, {"item_id": mx, "qty": 2, "measure": 0.5},
                                {"item_id": mk, "qty": 1, "measure": 1.5}, {"item_id": pf, "qty": 3}]})
    assert code == 200 and d["ok"], d
    db = w["db"]()
    lines = db.execute("SELECT name, qty, price, unit, measure, unit_price FROM cafe_bill_items ORDER BY id").fetchall()
    assert [tuple(r) for r in lines] == [("Mixture", 1, 12000, "kg", 0.25, 48000), ("Mixture", 2, 24000, "kg", 0.5, 48000),
                                         ("Milk", 1, 9000, "l", 1.5, 6000), ("Puffs", 3, 2000, "pc", None, None)]
    # 120 + 480 + 90 + 60 = 750
    assert db.execute("SELECT total FROM cafe_bills").fetchone()[0] == 75000
    page = b.get(d["print"]).text
    assert "250 g" in page and "1.5 L" in page and "₹480/kg" in page
    assert "₹480 / kg" in b.get("/cafe/menu").text or "/ kg" in b.get("/cafe/menu").text
    # analysis shows kg sold
    assert "1.25 kg" in b.get("/cafe/analysis").text
    # a weighed item without a weight counts as 1 kg; KOT shows the weight
    code, d = pos(b, {"action": "kot", "order_type": "takeaway", "items": [{"item_id": mx, "qty": 1, "measure": 0.75}]})
    kot = json.loads(w["db"]().execute("SELECT lines FROM cafe_kots ORDER BY id DESC").fetchone()[0])
    assert kot == [{"name": "Mixture", "qty": 1, "note": "", "measure": "750 g"}]


def test_price_change_needs_permission_and_open_order(cw):
    w = cw
    b = w["cafe"]
    settings(b)
    add_item(b, "Cake", "100")
    cake = item_id(w, "Cake")
    cashier = make_staff(b, "cash1")                                    # cashier preset: no price change
    code, d = pos(cashier, {"action": "bill", "order_type": "takeaway", "payment_mode": "Cash",
                            "items": [{"item_id": cake, "qty": 1, "price": "80"}]})
    assert code == 400 and "not allowed to change item prices" in d["error"]
    assert '"price_edit": false' in cashier.get("/cafe/pos").text
    # hold an order, then change its price (decrease and increase) with permission
    code, d = pos(b, {"action": "save", "order_type": "takeaway", "items": [{"item_id": cake, "qty": 2}]})
    assert d["ok"], d
    db = w["db"]()
    order = db.execute("SELECT id FROM cafe_orders WHERE status = 'open'").fetchone()[0]
    line = db.execute("SELECT id FROM cafe_order_items WHERE order_id = ?", (order,)).fetchone()[0]
    # cashier cannot change the held order's price either
    code, d = pos(cashier, {"action": "save", "order_id": order, "order_type": "takeaway",
                            "items": [{"line_id": line, "item_id": cake, "qty": 2, "price": "90"}]})
    assert code == 400
    code, d = pos(b, {"action": "save", "order_id": order, "order_type": "takeaway",
                      "items": [{"line_id": line, "item_id": cake, "qty": 2, "price": "120"}]})
    assert d["ok"], d
    row = w["db"]().execute("SELECT price, price_by FROM cafe_order_items WHERE id = ?", (line,)).fetchone()
    assert row["price"] == 12000 and row["price_by"]
    # the cashier can still bill it (price unchanged by them)
    code, d = pos(cashier, {"action": "bill", "order_id": order, "order_type": "takeaway", "payment_mode": "Cash",
                            "items": [{"line_id": line, "item_id": cake, "qty": 2}]})
    assert d["ok"], d
    db = w["db"]()
    assert db.execute("SELECT total FROM cafe_bills").fetchone()[0] == 24000
    assert db.execute("SELECT price_by FROM cafe_bill_items").fetchone()[0]
    assert db.execute("SELECT price FROM menu_items WHERE id = ?", (cake,)).fetchone()[0] == 10000  # menu unchanged
    assert "✏️" in b.get(d["view"]).text
    # a custom staff member with the tick can change prices
    mgr = make_staff(b, "mgr1", role="custom", perms=["bill_create", "price_edit"])
    code, d = pos(mgr, {"action": "bill", "order_type": "takeaway", "payment_mode": "Cash",
                        "items": [{"item_id": cake, "qty": 1, "price": "75.50"}]})
    assert d["ok"], d
    # customers can never send a price
    settings(b, online_on="1")
    r = Browser().post("/m/chai/order", {"payload": json.dumps({"type": "takeaway", "name": "A", "phone": "9876543210",
                       "items": [{"item_id": cake, "qty": 1, "price": "1"}]})}, follow=False)
    assert r.status_code == 400


def test_barcode_unique_and_in_excel(cw):
    w = cw
    b = w["cafe"]
    settings(b)
    add_item(b, "Biscuit", "10", barcode="8901234567890")
    r = b.post("/cafe/menu/item/new", {"name": "Other", "price": "5", "available": "1", "barcode": "8901234567890"})
    assert "already" in r.text and w["db"]().execute("SELECT COUNT(*) FROM menu_items").fetchone()[0] == 1
    assert '"barcode": "8901234567890"' in b.get("/cafe/pos").text
    assert "8901234567890" not in Browser().get("/m/chai").text            # not on the customer menu
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(b.get("/cafe/menu/export.xlsx").data))["Menu"]
    head = [c.value for c in ws[1]]
    assert "Barcode" in head and "Unit" in head
    assert ws.cell(2, head.index("Barcode") + 1).value == "8901234567890"
    sheet = xlsx([["S.No", "Category", "Item name", "Price", "Unit", "Barcode"],
                  [None, "Bakery", "Halwa", 400, "Kg", "123456"],
                  [None, "Bakery", "Rusk", 40, "Piece", "8901234567890"]])      # duplicate barcode -> error
    r = b.post("/cafe/menu/import", {"sheet": (sheet, "m.xlsx")}, content_type="multipart/form-data")
    assert "Import preview" in r.text
    tok = re.search(r'name="tok" value="([0-9a-f]+)"', r.text).group(1)
    b.post("/cafe/menu/import/confirm", {"tok": tok})
    db = w["db"]()
    halwa = db.execute("SELECT unit, barcode FROM menu_items WHERE name = 'Halwa'").fetchone()
    assert tuple(halwa) == ("kg", "123456")
    assert db.execute("SELECT COUNT(*) FROM menu_items WHERE name = 'Rusk'").fetchone()[0] == 0


def test_most_sold_first_and_pin(cw):
    w = cw
    b = w["cafe"]
    settings(b)
    for n in ("Aaa", "Bbb", "Ccc"):
        add_item(b, n, "10")
    a, bb, c = (item_id(w, n) for n in ("Aaa", "Bbb", "Ccc"))
    pos(b, {"action": "bill", "order_type": "takeaway", "payment_mode": "Cash", "items": [{"item_id": c, "qty": 5}, {"item_id": bb, "qty": 2}]})

    def order(page):
        names = re.findall(r'"name": "(Aaa|Bbb|Ccc)"', page)
        return names
    assert order(b.get("/cafe/pos").text) == ["Ccc", "Bbb", "Aaa"]
    assert order(Browser().get("/m/chai").text) == ["Ccc", "Bbb", "Aaa"]
    # pin Aaa to the top
    b.post(f"/cafe/menu/item/{a}/edit", {"name": "Aaa", "price": "10", "available": "1", "pinned": "1", "unit": "pc"})
    assert order(b.get("/cafe/pos").text) == ["Aaa", "Ccc", "Bbb"]
    # menu order setting for billing only
    settings(b, sort_pos="menu", sort_menu="popular")
    assert order(b.get("/cafe/pos").text)[0] == "Aaa"
    assert order(b.get("/cafe/pos").text)[1:] == ["Bbb", "Ccc"]
    assert order(Browser().get("/m/chai").text) == ["Aaa", "Ccc", "Bbb"]


def test_bluetooth_print_link(cw):
    w = cw
    b = w["cafe"]
    settings(b)
    add_item(b, "Tea", "15")
    code, d = pos(b, {"action": "bill", "order_type": "takeaway", "payment_mode": "Cash", "items": [{"item_id": item_id(w, "Tea"), "qty": 2}]})
    assert "intent:base64," not in b.get(d["print"]).text
    settings(b, bt_print="1")
    page = b.get(d["print"]).text
    m = re.search(r'intent:base64,([A-Za-z0-9+/=]+)#Intent;scheme=rawbt', page)
    assert m
    import base64
    raw = base64.b64decode(m.group(1))
    assert raw.startswith(b"\x1b@") and b"Tea" in raw and b"Rs.30" in raw.replace(b" ", b"")
    assert "intent:base64," in b.get(d["view"]).text


def test_public_order_with_weight(cw):
    w = cw
    b = w["cafe"]
    settings(b, online_on="1")
    add_item(b, "Mixture", "480", unit="kg")
    mx = item_id(w, "Mixture")
    assert '"unit": "kg"' in Browser().get("/m/chai").text
    r = Browser().post("/m/chai/order", {"payload": json.dumps({"type": "takeaway", "name": "Anu", "phone": "9876543210",
                       "items": [{"item_id": mx, "qty": 1, "measure": 0.5}]})}, follow=False)
    assert r.status_code == 200 and r.get_json()["ok"]
    row = w["db"]().execute("SELECT price, measure FROM cafe_order_items").fetchone()
    assert tuple(row) == (24000, 0.5)


# ------------------------------------------------------------------ round 4: speed, split / later payments, My Card

def test_gzip_and_static_cache(cw):
    b = cw["cafe"]
    settings(b)
    r = b.c.get("/cafe/pos", headers={"Accept-Encoding": "gzip"})
    assert r.headers.get("Content-Encoding") == "gzip"
    import gzip
    assert b"pos-cfg" in gzip.decompress(r.data)
    assert b.get("/cafe/pos").headers.get("Content-Encoding") is None      # browsers without gzip still work
    page = b.get("/cafe/pos").text
    assert re.search(r'/static/style\.css\?v=\d+', page)
    st = b.c.get(re.search(r'(/static/style\.css\?v=\d+)', page).group(1))
    assert "max-age=31536000" in st.headers["Cache-Control"]


def test_text_only_tiles(cw):
    w = cw
    b = w["cafe"]
    settings(b)
    b.post("/cafe/menu/item/new", {"name": "Cake", "price": "50", "available": "1", "photo": (png(), "c.png")},
           content_type="multipart/form-data")
    assert '"photo": "/media/' in b.get("/cafe/pos").text
    settings(b, pos_tiles="text")
    assert '"photo": null' in b.get("/cafe/pos").text
    assert "/media/" in Browser().get("/m/chai").text                      # customer menu keeps photos


def test_retry_does_not_make_two_bills(cw):
    w = cw
    b = w["cafe"]
    settings(b)
    add_item(b, "Tea", "15")
    p = {"action": "bill", "order_type": "takeaway", "payment_mode": "Cash", "client_ref": "abc123",
         "items": [{"item_id": item_id(w, "Tea"), "qty": 1}]}
    c1, d1 = pos(b, p)
    c2, d2 = pos(b, p)
    assert d1["ok"] and d1 == d2 and "embed=1" in d1["embed"]
    assert w["db"]().execute("SELECT COUNT(*) FROM cafe_bills").fetchone()[0] == 1
    # the in-screen print can be framed by our own page only
    r = b.get(d1["embed"])
    assert r.headers["X-Frame-Options"] == "SAMEORIGIN" and "pos-new" in r.text
    assert b.get(d1["print"]).headers["X-Frame-Options"] == "DENY"


def test_split_payment_and_reports(cw):
    w = cw
    b = w["cafe"]
    settings(b)
    add_item(b, "Meal", "350")
    meal = item_id(w, "Meal")
    base = {"action": "bill", "order_type": "takeaway", "items": [{"item_id": meal, "qty": 1}]}
    code, d = pos(b, dict(base, payment_mode="Split", payment_split={"cash": "200", "upi": "100", "card": ""}))
    assert code == 400 and "must be ₹350" in d["error"]
    code, d = pos(b, dict(base, payment_mode="Split", payment_split={"cash": "200", "upi": "150", "card": ""}))
    assert d["ok"], d
    bill = w["db"]().execute("SELECT * FROM cafe_bills").fetchone()
    assert (bill["payment_mode"], bill["pay_cash"], bill["pay_upi"], bill["pay_card"], bill["due"]) == ("Split", 20000, 15000, 0, 0)
    assert "Cash ₹200 + UPI ₹150" in b.get(d["print"]).text
    pos(b, dict(base, payment_mode="UPI"))
    page = b.get("/cafe/bills").text
    assert "₹200" in page and "₹500" in page                               # cash 200, UPI 150 + 350
    assert "Split" in b.get("/cafe/bills?mode=Split").text
    ana = b.get("/cafe/analysis").text
    assert "Cash" in ana and "UPI" in ana
    csv_ = b.get("/cafe/export/bills.csv?p=today").text
    assert "Cash,UPI,Card,Not paid" in csv_ and "Split" in csv_


def test_pay_later_receive_and_change_payment(cw):
    w = cw
    b = w["cafe"]
    settings(b)
    add_item(b, "Biriyani", "180")
    by = item_id(w, "Biriyani")
    code, d = pos(b, {"action": "bill", "order_type": "delivery", "payment_mode": "Later", "items": [{"item_id": by, "qty": 2}]})
    assert d["ok"], d
    db = w["db"]()
    bid = db.execute("SELECT id FROM cafe_bills").fetchone()[0]
    assert db.execute("SELECT due, pay_cash FROM cafe_bills").fetchone()[:] == (36000, 0)
    assert "NOT PAID" in b.get(d["print"]).text
    page = b.get("/cafe/bills").text
    assert "1 bill not paid: ₹360" in page
    assert "Receive payment" in b.get(f"/cafe/bills/{bid}").text
    # delivery boy brings ₹200 cash + ₹160 UPI
    r = b.post(f"/cafe/bills/{bid}/payment", {"mode": "Split", "cash": "200", "upi": "160", "card": ""})
    assert "Payment received: Cash ₹200 + UPI ₹160" in r.text
    row = w["db"]().execute("SELECT * FROM cafe_bills").fetchone()
    assert (row["due"], row["pay_cash"], row["pay_upi"], row["total"]) == (0, 20000, 16000, 36000)
    log = json.loads(row["pay_log"])
    assert log[0]["from"] == "Not paid ₹360" and log[0]["to"] == "Cash ₹200 + UPI ₹160"
    assert "bill not paid:" not in b.get("/cafe/bills").text
    # wrong total is refused
    r = b.post(f"/cafe/bills/{bid}/payment", {"mode": "Split", "cash": "100", "upi": "100"})
    assert "must be ₹360" in r.text
    # cashier: same day ok, an old bill not (without the tick); old unpaid bill ok
    cashier = make_staff(b, "cash5")
    r = cashier.post(f"/cafe/bills/{bid}/payment", {"mode": "UPI"})
    assert "Payment changed: UPI" in r.text
    with sqlite3.connect(w["db_path"]) as c:
        c.execute("UPDATE cafe_bills SET bill_date = '2020-01-01'")
    r = cashier.post(f"/cafe/bills/{bid}/payment", {"mode": "Cash"})
    assert "not allowed to change the payment" in r.text
    assert b.post(f"/cafe/bills/{bid}/payment", {"mode": "Cash"}).text.count("Payment changed: Cash") == 1   # admin
    with sqlite3.connect(w["db_path"]) as c:
        c.execute("UPDATE cafe_bills SET due = total, pay_cash = 0, payment_mode = 'Later'")
    assert "Payment received: Cash" in cashier.post(f"/cafe/bills/{bid}/payment", {"mode": "Cash"}).text
    # another cafe cannot touch it
    assert w["cafe2"].post(f"/cafe/bills/{bid}/payment", {"mode": "UPI"}, follow=False).status_code == 404


def test_my_card_page(cw):
    w = cw
    b = w["cafe"]
    settings(b, loyalty_on="1", referral_on="1", referral_referrer_points="50", referral_new_points="25", birthday_on="1")
    add_item(b, "Tea", "100")
    tea = item_id(w, "Tea")
    pos(b, {"action": "bill", "order_type": "takeaway", "payment_mode": "Cash", "items": [{"item_id": tea, "qty": 2}],
            "customer": {"phone": "9876543210", "name": "Anu"}})
    db = w["db"]()
    c = db.execute("SELECT * FROM customers").fetchone()
    assert c["card_token"] and len(c["card_token"]) >= 12
    url = f"/m/chai/me/{c['card_token']}"
    guest = Browser()
    r = guest.get(url)
    assert r.status_code == 200 and "Hi Anu" in r.text and ">10<" in r.text and "Your bills" in r.text
    assert c["referral_code"] in r.text and "Share on WhatsApp" in r.text
    assert "9876543210" not in r.text                                      # phone not shown
    # birthday from the card page
    guest.post(url + "/birthday", {"birthday": "1990-05-17"})
    assert w["db"]().execute("SELECT birthday FROM customers").fetchone()[0] == "05-17"
    # wrong token / other cafe -> 404
    assert guest.get("/m/chai/me/nottherealtoken123").status_code == 404
    assert guest.get(f"/m/other/me/{c['card_token']}").status_code == 404
    # link in the WhatsApp bill text and on the customer page; new link kills the old one
    bid = db.execute("SELECT id FROM cafe_bills").fetchone()[0]
    assert "me/" + c["card_token"] in b.get(f"/cafe/bills/{bid}").text
    assert "Send card link" in b.get(f"/cafe/customers/{c['id']}").text
    b.post(f"/cafe/customers/{c['id']}/new-card-link", {})
    assert guest.get(url).status_code == 404
    # print QR only when switched on
    assert 'class="c card-qr"' not in b.get(f"/cafe/bills/{bid}/print").text
    settings(b, loyalty_on="1", card_qr="1")
    assert "<svg" in b.get(f"/cafe/bills/{bid}/print").text
    # {my_card} in WhatsApp templates
    from cafe_people import fill_template
    import app as crm_app
    with crm_app.app.test_request_context():
        row = sqlite3.connect(w["db_path"])
        row.row_factory = sqlite3.Row
        cust = row.execute("SELECT * FROM customers").fetchone()
        org = row.execute("SELECT * FROM organizations WHERE code = 'CHAI'").fetchone()
        st = cafe.get_settings(org["id"])
        assert "/m/chai/me/" in fill_template("See {my_card}", cust, st, org)


def test_migration_v11_fills_old_bills_and_card_links(cw):
    w = cw
    b = w["cafe"]
    settings(b)
    add_item(b, "Tea", "15")
    pos(b, {"action": "bill", "order_type": "takeaway", "payment_mode": "UPI", "items": [{"item_id": item_id(w, "Tea"), "qty": 2}],
            "customer": {"phone": "9876500000", "name": "Old"}})
    with sqlite3.connect(w["db_path"]) as c:   # make it look like a bill/customer from before v11
        c.execute("UPDATE cafe_bills SET pay_upi = 0")
        c.execute("DROP INDEX uq_customers_card")
        c.execute("UPDATE customers SET card_token = NULL")
        c.execute("PRAGMA user_version = 10")
    crm_db.migrate(w["db_path"])
    db = w["db"]()
    assert db.execute("SELECT pay_upi, total FROM cafe_bills").fetchone()[:] == (3000, 3000)
    assert db.execute("SELECT card_token FROM customers").fetchone()[0]
    assert db.execute("PRAGMA user_version").fetchone()[0] == crm_db.SCHEMA_VERSION


# ------------------------------------------------------------------ round 5: order alerts, customer order retries

def test_public_order_old_token_and_retry(cw):
    w = cw
    b = w["cafe"]
    settings(b, online_on="1")
    add_item(b, "Tea", "15")
    tea = item_id(w, "Tea")
    guest = Browser()
    guest.get("/m/chai")
    payload = {"client_ref": "ref1", "type": "takeaway", "name": "Anu", "phone": "9876543210",
               "items": [{"item_id": tea, "qty": 1}]}
    # an old page: wrong token -> JSON with a fresh token (not an HTML error page)
    r = guest.post("/m/chai/order", {"payload": json.dumps(payload), "csrf_token": "old"}, follow=False)
    d = r.get_json()
    assert r.status_code == 400 and d["csrf"] and d["token"] == guest.token()
    # sent again with the fresh token -> ok; the same ref again (network retry) -> same answer, one order only
    r1 = guest.post("/m/chai/order", {"payload": json.dumps(payload), "csrf_token": d["token"]}, follow=False).get_json()
    r2 = guest.post("/m/chai/order", {"payload": json.dumps(payload)}, follow=False).get_json()
    assert r1["ok"] and r1 == r2
    assert w["db"]().execute("SELECT COUNT(*) FROM cafe_orders").fetchone()[0] == 1
    # a refused order can be fixed and sent with the same ref
    bad = dict(payload, client_ref="ref2", phone="12")
    assert guest.post("/m/chai/order", {"payload": json.dumps(bad)}, follow=False).status_code == 400
    good = dict(payload, client_ref="ref2", phone="9876500001")
    assert guest.post("/m/chai/order", {"payload": json.dumps(good)}, follow=False).get_json()["ok"]
    # other forms still get the normal error page
    assert b.post("/cafe/settings", {"name": "x", "csrf_token": "bad"}, follow=False).status_code == 400


def test_push_notifications_for_new_orders(cw, monkeypatch):
    import cafe_push
    sent = []

    class Gone(Exception):
        class response:
            status_code = 410

    def fake(sub, data, key):
        sent.append((sub["endpoint"], json.loads(data)))
        if "dead" in sub["endpoint"]:
            raise Gone()
    monkeypatch.setattr(cafe_push, "sender", fake)
    w = cw
    b = w["cafe"]
    settings(b, online_on="1")
    add_item(b, "Tea", "15")
    assert b.get("/sw.js").status_code == 200 and "showNotification" in b.get("/sw.js").text
    key = b.get("/cafe/push/key").get_json()["key"]
    assert len(key) > 80
    sub = {"endpoint": "https://fcm.googleapis.com/fcm/send/admin1", "keys": {"p256dh": "p", "auth": "a"}}
    assert b.post("/cafe/push/subscribe", {"payload": json.dumps(sub)}).get_json()["ok"]
    dead = dict(sub, endpoint="https://fcm.googleapis.com/fcm/send/dead1")
    b.post("/cafe/push/subscribe", {"payload": json.dumps(dead)})
    assert b.post("/cafe/push/subscribe", {"payload": json.dumps({"endpoint": "http://x"})}).status_code == 400
    # staff without online orders: cannot subscribe
    kitchen = make_staff(b, "kit1", role="custom", perms=["bill_create"])
    assert kitchen.post("/cafe/push/subscribe", {"payload": json.dumps(dict(sub, endpoint="https://e/k"))},
                        follow=False).status_code == 403
    # another cafe's phone never gets this cafe's orders
    w["cafe2"].post("/cafe/push/subscribe", {"payload": json.dumps(dict(sub, endpoint="https://e/other"))})
    guest = Browser()
    guest.get("/m/chai")
    r = guest.post("/m/chai/order", {"payload": json.dumps({"type": "takeaway", "name": "Anu", "phone": "9876543210",
                                                            "items": [{"item_id": item_id(w, "Tea"), "qty": 2}]})},
                   follow=False)
    assert r.get_json()["ok"]
    endpoints = sorted(e for e, _ in sent)
    assert endpoints == ["https://fcm.googleapis.com/fcm/send/admin1", "https://fcm.googleapis.com/fcm/send/dead1"]
    msg = sent[0][1]
    assert "New online order #1" in msg["title"] and "Anu" in msg["body"] and "₹30" in msg["body"]
    left = [r[0] for r in w["db"]().execute("SELECT endpoint FROM push_subscriptions WHERE org_id = ?", (w["oid"]("CHAI"),))]
    assert left == ["https://fcm.googleapis.com/fcm/send/admin1"]      # the gone phone was removed
    # test button and turning off
    sent.clear()
    assert b.post("/cafe/push/test", {}).get_json()["ok"] and len(sent) == 1
    b.post("/cafe/push/unsubscribe", {"endpoint": "https://fcm.googleapis.com/fcm/send/admin1"})
    assert not b.post("/cafe/push/test", {}).get_json()["ok"]


def test_orders_page_alert_tools(cw):
    b = cw["cafe"]
    settings(b, online_on="1")
    page = b.get("/cafe/orders").text
    assert 'id="orders-live"' in page and "push-btn" in page and "awake-btn" in page
    base = b.get("/cafe/pos").text
    assert "setInterval(function () { if (count) bell(); }, 3000)" in base


def test_home_screen_app_files(cw):
    b = cw["cafe"]
    page = b.get("/cafe/orders").text
    assert "/chai/app.webmanifest" in page and "apple-touch-icon" in page and "sound-btn" in page
    data = json.loads(b.get("/chai/app.webmanifest").data)
    assert data["display"] == "standalone" and len(data["icons"]) == 2 and data["name"] == "Chai Point"
    assert data["start_url"] == "/chai"
    assert b.get("/static/icon-512.png").status_code == 200
    assert "manifest.webmanifest" in Browser().get("/login").text      # Real View app before login


def test_push_reminders_until_handled(cw, monkeypatch):
    import cafe_push
    sent = []
    monkeypatch.setattr(cafe_push, "sender", lambda sub, data, key: sent.append(json.loads(data)))
    cafe_push._last_sent.clear()
    w = cw
    b = w["cafe"]
    settings(b, online_on="1")
    add_item(b, "Tea", "15")
    b.post("/cafe/push/subscribe", {"payload": json.dumps({"endpoint": "https://fcm.googleapis.com/x/1",
                                                           "keys": {"p256dh": "p", "auth": "a"}})})
    guest = Browser()
    guest.get("/m/chai")
    guest.post("/m/chai/order", {"payload": json.dumps({"type": "takeaway", "name": "Anu", "phone": "9876543210",
                                                        "items": [{"item_id": item_id(w, "Tea"), "qty": 1}]})})
    assert len(sent) == 1 and "New online order" in sent[0]["title"]
    import app as crm_app
    import time as _t
    with crm_app.app.app_context():
        assert cafe_push.remind_waiting(_t.time() + 5) == 0          # too soon
        assert cafe_push.remind_waiting(_t.time() + 31) == 1         # 30 s later: reminder
        assert "still waiting" in sent[-1]["title"]
    oid_ = w["db"]().execute("SELECT id FROM cafe_orders").fetchone()[0]
    b.post(f"/cafe/orders/{oid_}/confirm", {})
    with crm_app.app.app_context():
        assert cafe_push.remind_waiting(_t.time() + 100) == 0        # handled: no more reminders
    # old waiting orders (over 30 minutes) stop ringing
    guest.post("/m/chai/order", {"payload": json.dumps({"type": "takeaway", "name": "Anu", "phone": "9876543211",
                                                        "items": [{"item_id": item_id(w, "Tea"), "qty": 1}]})})
    with sqlite3.connect(w["db_path"]) as c:
        c.execute("UPDATE cafe_orders SET created_at = '2020-01-01 10:00:00' WHERE status = 'pending'")
    with crm_app.app.app_context():
        assert cafe_push.remind_waiting(_t.time() + 200) == 0
